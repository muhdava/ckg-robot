r"""
Template DATA SISWA untuk diisi OPERATOR SEKOLAH.

    python alat\buat_template_sekolah.py

Menulis template\Template DATA SISWA (untuk sekolah).xlsx.

Kenapa template sendiri. Kerugian terbesar CKG Sekolah datang dari NIK yang
rusak SEBELUM sampai ke puskesmas: Excel menyimpan NIK sebagai ANGKA dan
membuang digit ke-16 (Al Irsyad 83 dari 173, SMPN 136 31 dari 600). Digit yang
hilang tidak bisa dikembalikan. Template ini mencegahnya di tempat data diketik:
  - kolom NIK berformat TEKS, dengan validasi 16 angka saat diketik;
  - kolom "Cek NIK" menjelaskan masalahnya dalam kalimat (angka, bukan 16 digit,
    terpotong 0000, kembar, tanggal lahir/jenis kelamin tidak cocok dengan NIK);
  - baris bermasalah berwarna, supaya terlihat walau NIK-nya DITEMPEL (validasi
    Excel tidak memeriksa data yang ditempel);
  - lembar Petunjuk untuk operator sekolah.
Judul kolomnya dikenali konverter aplikasi (Siapkan data dari Excel) apa adanya;
kolom "Cek NIK" dilewati.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AKAR = Path(__file__).resolve().parent.parent
TUJUAN = AKAR / "template" / "Template DATA SISWA (untuk sekolah).xlsx"
AWAL, AKHIR = 5, 1004                          # 1000 murid per berkas

KOLOM = [  # (judul, lebar, format, keterangan singkat di baris 3)
    ("No", 5, None, "otomatis"),
    ("NIK", 19, "@", "16 angka, dari KK"),
    ("Cek NIK", 34, None, "otomatis - harus ✓ OK"),
    ("Nama", 30, None, "sesuai KK"),
    ("Tanggal Lahir", 13, "DD/MM/YYYY", "16/05/2012"),
    ("Jenis Kelamin", 13, None, "pilih"),
    ("Alamat", 26, None, "jalan / RT RW / kelurahan"),
    ("Nama Sekolah", 26, None, "sama dengan portal"),
    ("Kelas", 9, None, "pilih"),
    ("Berat Badan (kg)", 10, "0.0", "diisi puskesmas"),
    ("Tinggi Badan (cm)", 10, "0.0", "diisi puskesmas"),
    ("Hemoglobin (g/dL)", 11, "0.0", "diisi puskesmas"),
    ("Skrining Gigi", 12, None, "diisi puskesmas"),
    ("Serumen", 10, None, "diisi puskesmas"),
    ("Visus Mata", 12, None, "diisi puskesmas"),
    ("Kebugaran", 12, None, "diisi puskesmas"),
]
HIJAU, MERAH_MUDA, JINGGA_MUDA = "0F766E", "FDE2DF", "FFF1D6"


def cek_nik(r: int) -> str:
    """Rumus kolom C: kalimat yang menjelaskan masalah NIK, atau ✓ OK."""
    b, e, f = f"B{r}", f"E{r}", f"F{r}"
    hari = f"VALUE(MID({b},7,2))"
    return (
        f'=IF({b}="","",'
        f'IF(NOT(ISTEXT({b})),"✗ NIK tersimpan sebagai ANGKA - hapus, lalu ketik ulang dari KK",'
        f'IF(OR(LEN({b})<>16,ISERROR(--{b})),"✗ NIK harus tepat 16 angka",'
        f'IF(RIGHT({b},4)="0000","✗ NIK terpotong (berakhir 0000) - ketik ulang dari KK",'
        f'IF(SUMPRODUCT(--($B${AWAL}:$B${AKHIR}={b}))>1,"✗ NIK kembar dengan murid lain",'
        f'IF(AND(ISNUMBER({e}),OR({hari}-IF({hari}>40,40,0)<>DAY({e}),VALUE(MID({b},9,2))<>MONTH({e}),'
        f'VALUE(MID({b},11,2))<>MOD(YEAR({e}),100))),"⚠ tanggal lahir tidak cocok dengan NIK",'
        f'IF(AND({f}<>"",({hari}>40)<>({f}="Perempuan")),"⚠ jenis kelamin tidak cocok dengan NIK",'
        f'"✓ OK")))))))'
    )


def daftar(ws, sqref: str, isi: str, judul: str, pesan: str, stop: bool = True) -> None:
    dv = DataValidation(type="list", formula1=f'"{isi}"', allow_blank=True, showErrorMessage=True,
                        errorStyle="stop" if stop else "warning", errorTitle=judul, error=pesan)
    dv.add(sqref)
    ws.add_data_validation(dv)


def angka(ws, sqref: str, bawah: float, atas: float, judul: str, satuan: str) -> None:
    dv = DataValidation(type="decimal", operator="between", formula1=str(bawah), formula2=str(atas),
                        allow_blank=True, showErrorMessage=True, errorStyle="warning", errorTitle=judul,
                        error=f"{judul} biasanya {bawah}-{atas} {satuan}. Periksa lagi angkanya "
                              "(pakai koma/titik desimal sesuai Excel).")
    dv.add(sqref)
    ws.add_data_validation(dv)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Siswa"
    gelap = PatternFill("solid", fgColor=HIJAU)
    putih = Font(bold=True, color="FFFFFF")
    tipis = Side(style="thin", color="D0D7DE")

    ws["A1"] = "DATA SISWA - Cek Kesehatan Gratis (CKG) Sekolah"
    ws["A1"].font = Font(bold=True, size=14, color=HIJAU)
    ws["A2"] = ("Isi mulai baris 5, satu murid per baris. Kolom NIK diketik dari Kartu Keluarga (jangan "
                "ditempel dari berkas yang NIK-nya tampil 3,17E+15). Kolom 'Cek NIK' harus menunjukkan ✓ OK. "
                "Kolom berwarna hijau muda diisi petugas puskesmas saat pemeriksaan. Baca lembar 'Petunjuk'.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:P2")
    ws.row_dimensions[2].height = 42
    for k, (judul, lebar, fmt, ket) in enumerate(KOLOM, 1):
        h = ws.cell(4, k, judul)
        h.font, h.fill = putih, gelap
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = Border(bottom=tipis)
        c3 = ws.cell(3, k, ket)
        c3.font = Font(italic=True, size=8, color="64748B")
        c3.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[h.column_letter].width = lebar
    ws.row_dimensions[4].height = 30
    periksa = PatternFill("solid", fgColor="EEF7F3")
    for r in range(AWAL, AKHIR + 1):
        ws.cell(r, 1, f'=IF(B{r}="","",ROW()-{AWAL - 1})').font = Font(color="94A3B8")
        ws.cell(r, 3, cek_nik(r))
        for k, (_, _, fmt, _) in enumerate(KOLOM, 1):
            if fmt:
                ws.cell(r, k).number_format = fmt
        for k in range(10, 17):
            ws.cell(r, k).fill = periksa
    ws.freeze_panes = "D5"

    rng = f"{AWAL}:{AKHIR}"
    # NIK: diketik harus teks 16 angka (validasi Excel hanya berlaku saat MENGETIK).
    dv = DataValidation(type="custom", formula1=f"AND(ISTEXT(B{AWAL}),LEN(B{AWAL})=16,NOT(ISERROR(--B{AWAL})))",
                        allow_blank=True, showErrorMessage=True, errorStyle="stop", errorTitle="NIK",
                        error="NIK harus tepat 16 angka (lihat Kartu Keluarga). Jangan ada spasi atau titik.",
                        showInputMessage=True, promptTitle="NIK",
                        prompt="Ketik 16 angka NIK dari Kartu Keluarga. Jangan tempel dari berkas yang NIK-nya tampil 3,17E+15.")
    dv.add(f"B{AWAL}:B{AKHIR}")
    ws.add_data_validation(dv)
    dv = DataValidation(type="date", operator="between", formula1="DATE(1990,1,1)", formula2="TODAY()",
                        allow_blank=True, showErrorMessage=True, errorStyle="warning", errorTitle="Tanggal lahir",
                        error="Isi tanggal lahir seperti 16/05/2012 (hari/bulan/tahun).")
    dv.add(f"E{AWAL}:E{AKHIR}")
    ws.add_data_validation(dv)
    daftar(ws, f"F{AWAL}:F{AKHIR}", "Laki-laki,Perempuan", "Jenis kelamin", "Pilih Laki-laki atau Perempuan.")
    daftar(ws, f"I{AWAL}:I{AKHIR}", ",".join(f"Kelas {i}" for i in range(1, 13)), "Kelas",
           "Pilih dari daftar: Kelas 1 sampai Kelas 12.")
    angka(ws, f"J{AWAL}:J{AKHIR}", 8, 150, "Berat badan", "kg")
    angka(ws, f"K{AWAL}:K{AKHIR}", 80, 210, "Tinggi badan", "cm")
    angka(ws, f"L{AWAL}:L{AKHIR}", 4, 20, "Hemoglobin", "g/dL")
    daftar(ws, f"M{AWAL}:M{AKHIR}", "Normal,Sedang,Tidak Normal", "Skrining gigi", "Pilih dari daftar.")
    daftar(ws, f"N{AWAL}:N{AKHIR}", "Ada,Tidak Ada", "Serumen", "Pilih dari daftar.")
    daftar(ws, f"O{AWAL}:O{AKHIR}", "Normal,Visus Kanan,Visus Kiri,Tidak Normal", "Visus mata", "Pilih dari daftar.")
    daftar(ws, f"P{AWAL}:P{AKHIR}", "Baik sekali,Baik,Cukup,Kurang,Kurang sekali", "Kebugaran", "Pilih dari daftar.")

    # Warna: juga menangkap NIK yang DITEMPEL (validasi tidak memeriksa tempelan).
    ws.conditional_formatting.add(f"B{AWAL}:C{AKHIR}", FormulaRule(
        formula=[f'LEFT($C{AWAL},1)="✗"'], fill=PatternFill(fill_type="solid", start_color=MERAH_MUDA, end_color=MERAH_MUDA),
        font=Font(color="9B1C12", bold=True)))
    ws.conditional_formatting.add(f"B{AWAL}:C{AKHIR}", FormulaRule(
        formula=[f'LEFT($C{AWAL},1)="⚠"'], fill=PatternFill(fill_type="solid", start_color=JINGGA_MUDA, end_color=JINGGA_MUDA),
        font=Font(color="8A4B00")))
    ws.conditional_formatting.add(f"C{AWAL}:C{AKHIR}", FormulaRule(
        formula=[f'LEFT($C{AWAL},1)="✓"'], font=Font(color="15803D")))
    _ = rng

    # ------------------------------------------------------------ Petunjuk
    pt = wb.create_sheet("Petunjuk")
    pt.column_dimensions["A"].width = 4
    pt.column_dimensions["B"].width = 110
    isi = [
        ("judul", "Petunjuk mengisi DATA SISWA untuk CKG Sekolah"),
        ("", "Berkas ini dipakai puskesmas untuk mendaftarkan dan mencatat hasil Cek Kesehatan Gratis murid "
             "di portal Sehat Indonesiaku. Data yang salah membuat murid DITOLAK portal, atau lebih buruk: "
             "tercatat atas nama orang lain."),
        ("sub", "1. Yang diisi sekolah"),
        ("", "NIK, Nama, Tanggal Lahir, Jenis Kelamin, Alamat, Nama Sekolah, Kelas - satu murid per baris, mulai baris 5."),
        ("", "Kolom berlatar hijau muda (Berat Badan sampai Kebugaran) DIKOSONGKAN - diisi petugas puskesmas saat pemeriksaan."),
        ("sub", "2. NIK - bagian terpenting"),
        ("", "Ketik 16 angka NIK dari Kartu Keluarga. Kolomnya sudah berformat teks, jadi angkanya tidak akan berubah."),
        ("", "JANGAN menempel NIK dari berkas lain yang NIK-nya tampil seperti 3,17204E+15 atau berakhir 000. "
             "Di berkas itu digit terakhirnya SUDAH HILANG dan tidak bisa dikembalikan - ketik ulang dari KK."),
        ("", "Kalau terpaksa menempel: klik kanan → Paste Special → Values, lalu periksa kolom Cek NIK."),
        ("sub", "3. Kolom 'Cek NIK' (otomatis)"),
        ("", "✓ OK  - NIK siap dipakai."),
        ("", "✗ NIK tersimpan sebagai ANGKA - NIK ditempel sebagai angka. Hapus, lalu ketik ulang."),
        ("", "✗ NIK harus tepat 16 angka - kurang/lebih digit, atau ada spasi/titik."),
        ("", "✗ NIK terpotong (berakhir 0000) - digit terakhir hilang. Ketik ulang dari KK."),
        ("", "✗ NIK kembar dengan murid lain - dua murid ber-NIK sama; salah satunya salah ketik."),
        ("", "⚠ tanggal lahir / jenis kelamin tidak cocok dengan NIK - NIK memuat tanggal lahir (digit 7-12, "
             "perempuan tanggalnya +40). Periksa mana yang salah: NIK atau tanggal lahirnya."),
        ("", "Baris dengan ✗ berwarna merah, ⚠ berwarna jingga. Kirim berkas setelah semua baris ✓ OK."),
        ("sub", "4. Nama Sekolah dan Kelas"),
        ("", "Nama Sekolah ditulis SAMA PERSIS dengan nama di portal Sehat Indonesiaku (tanyakan ke puskesmas bila ragu), "
             "mis. 'SMP Negeri 136 Jakarta'. Kelas dipilih dari daftar: Kelas 1 sampai Kelas 12."),
        ("sub", "5. Setelah selesai"),
        ("", "Simpan berkas ini apa adanya (.xlsx) dan kirim ke petugas puskesmas. Tidak perlu diubah ke CSV."),
        ("", "Murid yang belum punya NIK atau NIK-nya belum diketahui tetap boleh ditulis - biarkan NIK kosong; "
             "petugas akan menindaklanjuti."),
    ]
    r = 1
    for jenis, teks in isi:
        c = pt.cell(r, 2, teks)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if jenis == "judul":
            c.font = Font(bold=True, size=14, color=HIJAU)
        elif jenis == "sub":
            c.font = Font(bold=True, size=11.5, color=HIJAU)
            r += 0
        pt.row_dimensions[r].height = 18 if jenis else max(18, 15 * (len(teks) // 105 + 1))
        r += 1 if jenis != "sub" else 1
    wb.active = 0
    TUJUAN.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TUJUAN)
    print(TUJUAN)


if __name__ == "__main__":
    main()
