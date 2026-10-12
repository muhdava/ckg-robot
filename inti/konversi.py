"""
Kenali jenis berkas mentah, lalu serahkan ke pengekstrak yang tepat.

    python inti\\konversi.py "C:\\...\\Data SDN LAGOA 09.xlsx"
    python inti\\konversi.py "...xlsx" --kelas 1 5

KENAPA DIKENALI, BUKAN DIPILIH SENDIRI
    Ada lima bentuk berkas yang beredar, dan memilih bentuk yang salah tidak
    menghasilkan galat - hanya CSV yang isinya melenceng. Sementara itu tiap
    bentuk punya penanda yang tidak bisa keliru: judul di sel A1, atau nama
    lembarnya. Jadi lebih aman ditanyakan pada berkasnya sendiri.

    Penanda di bawah dibaca dari 57 berkas nyata di folder Downloads pengguna,
    bukan dikarang: 30 register sekolah, 2 skrining TBC, 3 buku kunjungan, 1
    laporan SRQ, dan sisanya bentuk lain-lain.

BERKAS KELUARAN SENDIRI IKUT DIKENALI
    `hasil-*.xlsx` adalah HASIL macro, bukan sumber data. Ia punya kolom NIK dan
    Nama, jadi pengekstrak generik akan dengan senang hati "berhasil" membacanya
    dan menghasilkan daftar peserta palsu. Karena itu ia dikenali dan ditolak.
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import config                                     # noqa: E402
import ambil_data_umum as umum                    # noqa: E402


def _intip(path: Path) -> tuple[list[str], str]:
    """Nama-nama lembar dan teks beberapa baris pertama lembar pertama."""
    if path.suffix.lower() in (".csv", ".txt"):
        isi = path.read_text(encoding="utf-8-sig", errors="replace")
        return [path.stem], "\n".join(isi.splitlines()[:8])
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    lembar = list(wb.sheetnames)
    ws = wb[lembar[0]]
    baris = []
    for i, r in enumerate(ws.iter_rows(max_row=10, values_only=True)):
        baris.append(" | ".join(str(x) for x in r if x is not None))
    wb.close()
    return lembar, "\n".join(baris)


def kenali_sumber(path: Path) -> tuple[str, str]:
    """Kembalikan (jenis, alasan). Jenis 'umum' = tak ada bentuk khusus yang cocok."""
    try:
        lembar, atas = _intip(path)
    except Exception as e:
        return "gagal", f"berkas tidak bisa dibuka: {type(e).__name__}"
    nama_lembar = " | ".join(lembar)

    if re.search(r"REGISTER KEGIATAN KESEHATAN ANAK USIA SEKOLAH", atas, re.I):
        return "sekolah", "judul 'REGISTER KEGIATAN KESEHATAN ANAK USIA SEKOLAH'"
    if "Skrining" in lembar and "Rekap" in lembar:
        return "skrining_tb", "punya lembar 'Skrining' dan 'Rekap'"
    if re.search(r"MRN.*NAMA PASIEN|KUNJUNGAN RAWAT JALAN", atas, re.I | re.S):
        return "balita", "buku kunjungan rawat jalan (kolom MRN / NAMA PASIEN)"
    if re.search(r"Report\s+SRQ", atas, re.I):
        return "srq", "judul 'Report SRQ'"
    if re.match(r"^hasil-\d", path.stem, re.I) or re.search(
            r"^No \| NIK \| Nama \| Tanggal Lahir \| Status", atas):
        return "keluaran", "ini berkas HASIL macro, bukan data sumber"
    return "umum", f"tidak cocok bentuk khusus mana pun (lembar: {nama_lembar[:60]})"


SEBUTAN = {
    "sekolah":     "Register Kegiatan Kesehatan sekolah",
    "skrining_tb": "Lembar Skrining pemantauan TBC",
    "balita":      "Buku kunjungan rawat jalan (balita 0-6 th)",
    "srq":         "Laporan SRQ",
    "umum":        "Bentuk bebas - dipetakan dari judul kolom",
    "keluaran":    "Berkas hasil macro (bukan sumber)",
    "gagal":       "Tidak bisa dibaca",
}


def _tulis(baris: list[list], catatan: list[str], nama: str,
           laporan: dict) -> None:
    tujuan = config.DATA / f"datackg-{nama}.csv"
    with tujuan.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(baris)
    laporan.update(jumlah=len(baris), tujuan=str(tujuan))
    laporan["catatan"].extend(catatan)


def _tanda(path: Path, akhiran: str = "") -> str:
    dasar = re.sub(r"[^a-z0-9]+", "-",
                   re.sub(r"^data\s+", "", path.stem, flags=re.I).lower()).strip("-")
    return f"{dasar[:60]}{akhiran}"


# Lembar yang BUKAN data peserta walau memuat kolom NIK & Nama: laporan dan
# daftar periksa buatan alat perapian (rapikan_smp, laporan_skrining). Dulu
# converter selalu membaca lembar PERTAMA - dan berkas "RAPI" menaruh lembar
# Laporan di depan, sehingga tabel laporan diubah jadi 29 "peserta" palsu
# tanpa satu pun galat (SMPN 136, 6 Okt 2026).
_BUKAN_DATA = re.compile(r"laporan|ringkasan|rekap|perlu\s*dicek|panduan|petunjuk|referensi|penjelasan|alasan", re.I)


def lembar_data(path: Path) -> list[str]:
    """Nama lembar yang berisi data peserta: ada judul NIK dan Nama, bukan laporan."""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return []
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    hasil = []
    for ws in wb.worksheets:
        if _BUKAN_DATA.search(ws.title):
            continue
        for r in ws.iter_rows(max_row=40, values_only=True):
            judul = [str(x).strip().lower() for x in r if x is not None]
            if any(re.fullmatch(r"n\.?i\.?k\.?", j) for j in judul) and any(j.startswith("nama") for j in judul):
                hasil.append(ws.title)
                break
    wb.close()
    return hasil


def jalankan(path: Path, jenis: str = "", kelas: list[int] | None = None,
             sekolah: str = "", lembar: str = "", maks: int = 0,
             bentuk: str = "auto", ikut_nik: bool = True, tahan_kosong: bool = True) -> dict:
    """
    Ubah satu berkas mentah jadi CSV siap jalan.

    Tidak pernah melempar galat karena datanya cacat: yang cacat masuk ke
    `peringatan` / `catatan`, supaya satu berkas buruk tidak menjatuhkan
    aplikasi yang memanggilnya.
    """
    path = Path(path)
    laporan: dict = {"berkas": path.name, "jenis": "", "sebutan": "", "alasan": "",
                     "jumlah": 0, "tujuan": "", "peringatan": [], "catatan": []}
    if not path.exists():
        laporan["peringatan"].append(f"BERHENTI: berkas tidak ada: {path}")
        return laporan

    # Sasaran BALITA TANPA NIK (21), LANSIA parameter (33), dan LANSIA real
    # data (34) tidak punya pengekstrak khusus - hanya jalur generiklah yang
    # bisa menghasilkannya. Tanpa paksaan ini, berkas
    # yang kebetulan berbentuk buku kunjungan akan diserahkan ke pengekstrak
    # balita biasa dan keluar 18 kolom, padahal pengguna jelas-jelas memilih
    # sasaran yang lain di layar.
    if bentuk in ("balita_tanpa_nik", "lansia", "lansia_real"):
        jenis = "umum"
    jenis = jenis or kenali_sumber(path)[0]
    laporan["jenis"] = jenis
    laporan["sebutan"] = SEBUTAN.get(jenis, jenis)
    laporan["alasan"] = kenali_sumber(path)[1]

    if jenis == "gagal":
        laporan["peringatan"].append("BERHENTI: " + laporan["alasan"])
        return laporan
    if jenis == "keluaran":
        laporan["peringatan"].append(
            "BERHENTI: ini berkas HASIL macro, bukan data sumber. Mengonversinya "
            "akan menghasilkan daftar peserta palsu.")
        return laporan
    if jenis == "srq":
        laporan["peringatan"].append(
            "Laporan SRQ perlu rentang baris, jadi belum bisa lewat halaman ini. "
            "Jalankan: python inti\\ambil_data_srq.py \"<berkas>\" <awal> <akhir>")
        return laporan

    try:
        if jenis == "sekolah":
            import ambil_data_sekolah as m
            baris, catatan = m.olah(path, kelas or [], sekolah)
            akhiran = "-kelas" + "".join(str(k) for k in kelas) if kelas else ""
            if not baris:
                laporan["peringatan"].append(
                    "BERHENTI: tidak ada murid yang cocok. Periksa nomor kelasnya.")
                return laporan
            _tulis(baris, catatan, _tanda(path, akhiran), laporan)
            laporan["peringatan"].append(
                f"nama sekolah ditulis {baris[0][10]!r} - COCOKKAN dengan tulisan di portal.")
        elif jenis == "balita":
            import ambil_data_balita as m
            baris, catatan = m.olah(path, maks)
            if not baris:
                laporan["peringatan"].append(
                    "BERHENTI: tidak ada peserta berumur 0-6 tahun di berkas ini.")
                return laporan
            _tulis(baris, catatan, "balita-" + _tanda(path), laporan)
        elif jenis == "skrining_tb":
            import ambil_data_skrining_tb as m
            baris, catatan = m.olah(path, maks)
            if not baris:
                laporan["peringatan"].append("BERHENTI: lembar Skrining tidak berisi peserta.")
                return laporan
            _tulis(baris, catatan, _tanda(path), laporan)
        else:
            # Lembar mana yang dibaca: yang diminta, atau SEMUA lembar data
            # (berkas gabungan berisi satu lembar per sekolah -> satu CSV per
            # lembar). Lembar laporan/ringkasan dilewati.
            daftar = [lembar] if lembar else (lembar_data(path) or [None])
            if len(daftar) > 1:
                laporan["peringatan"].append(
                    f"{len(daftar)} lembar data, masing-masing jadi satu CSV: " + ", ".join(daftar))
            tujuan = []
            for nama_lembar in daftar:
                lap = umum.konversi(path, bentuk=bentuk or "auto",
                                    lembar=nama_lembar, sekolah=sekolah,
                                    kelas=(f"Kelas {kelas[0]}" if kelas else ""),
                                    ikut_nik=ikut_nik, tahan_kosong=tahan_kosong)
                awal = f"[{nama_lembar}] " if len(daftar) > 1 else ""
                laporan["jumlah"] += lap["jumlah"]
                laporan["ditahan"] = laporan.get("ditahan", 0) + lap.get("ditahan", 0)
                if lap["tujuan"]:
                    tujuan.append(lap["tujuan"])
                laporan["peringatan"].extend(awal + x for x in lap["peringatan"])
                laporan["catatan"].extend(awal + x for x in lap["catatan"])
            laporan["tujuan"] = ", ".join(tujuan)
    except Exception as e:
        laporan["peringatan"].append(
            f"BERHENTI: pengekstrak '{jenis}' gagal - {type(e).__name__}: "
            + " ".join(str(e).split())[:200])
    return laporan


def _cetak(l: dict) -> None:
    print(f"\nBerkas  : {l['berkas']}")
    print(f"Dikenali: {l['sebutan']}  ({l['alasan']})")
    print(f"Terbaca : {l['jumlah']} baris")
    if l["tujuan"]:
        print(f"Ditulis : {l['tujuan']}")
    for p in l["peringatan"]:
        print(f"  [PERINGATAN] {p}")
    for c in l["catatan"][:15]:
        print(f"  [CEK] {c}")
    if len(l["catatan"]) > 15:
        print(f"  [CEK] ... dan {len(l['catatan']) - 15} lagi")


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 1
    kelas = [int(x) for x in argv[1:] if x.isdigit()]
    ambil = lambda t: (argv[argv.index(t) + 1] if t in argv else "")   # noqa: E731
    l = jalankan(Path(argv[0]), kelas=kelas, sekolah=ambil("--sekolah"),
                 lembar=ambil("--lembar"))
    _cetak(l)
    return 0 if l["jumlah"] else 2


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
