"""
Pilah buku kunjungan puskesmas menjadi tiga CSV menurut umur.

    python inti\\pilah_kunjungan.py "C:\\...\\KUNJUNGAN JULI 2026 UNTUK CKG.xlsx"

    balita        0-6 tahun    -> tata letak BALITA (18 kolom)
    remaja        7-18 tahun   -> tata letak PARAMETER (33 kolom)
    dewasa/lansia 19 tahun +   -> tata letak PARAMETER (33 kolom)

BENTUK BERKAS SUMBER
    Kolomnya: No, Jam, MRN, Nama Pasien, L/P, "Tgl. Lahir / Umur", NIK,
    No. BPJS, Alamat, L/B. Baris kepalanya TIDAK selalu di tempat yang sama -
    berkas Juni memakai baris 4, berkas Juli baris 1 - jadi dicari, bukan
    dihafal.

    Umur dihitung dari TANGGAL LAHIR terhadap HARI INI, bukan dari angka
    "(11 Thn)" di dalam kurung: angka itu umur saat kunjungan, dan berkas
    kunjungan bulan lalu akan menaruh orang di kategori yang salah.

KENAPA REMAJA MEMAKAI TATA LETAK PARAMETER, BUKAN TATA LETAK SEKOLAH
    Sebelumnya berkas remaja ditulis dengan tata letak sekolah lama (25 kolom).
    Tata letak itu TIDAK PERNAH TERBACA sebagai sekolah, dan diamnya berbahaya:
    `peserta._geser_sekolah` mengenali berkas sekolah dari tulisan "Kelas N" di
    kolom 12 atau 16, sedangkan anak yang datang ke puskesmas tidak punya
    sekolah maupun kelas - kolomnya kosong. Berkasnya lalu jatuh ke tata letak
    PARAMETER, dan seluruh kolom sesudah wilayah bergeser: nilai hemoglobin
    terbaca sebagai hasil skrining mata dan telinga, serumen terbaca sebagai
    karies. Salah, tanpa satu pun pesan galat.

    Lagi pula alur sekolah memang bukan alurnya: murid dicari lewat Sekolah +
    Kelas, sedangkan anak-anak ini datang sendiri ke puskesmas dan didaftarkan
    lewat alur perorangan - alur yang justru butuh rantai wilayah, dan sudah
    menangani umur di bawah 18 (`pendaftaran._isi_wali`). Portal sendiri yang
    menentukan kumpulan formnya dari umur peserta.
"""
from __future__ import annotations

import csv
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import config                                                # noqa: E402
from peserta import bersihkan_nama                           # noqa: E402
from ambil_data_sekolah import dari_nik                       # noqa: E402
from ambil_data_balita import tanggal_kunjungan, umur_tahun  # noqa: E402

BALITA_MAKS, REMAJA_MAKS = 7, 19       # batas atas (eksklusif) tiap kategori

# Nilai tetap per tata letak. Kolomnya dikenali dari POSISI, jadi angka-angka
# di bawah ini harus sejalan dengan template masing-masing.
TETAP_BALITA = {6: "DKI Jakarta", 7: "Kota Adm. Jakarta Utara", 8: "Koja",
                9: "Lagoa", 11: "Belum Menikah", 12: "Tidak Bekerja",
                14: "Tidak", 15: "normal", 16: "normal", 17: "Ya", 18: "normal"}
# Remaja dan dewasa berbagi tata letak yang sama, jadi nilai tetapnya juga
# sama; yang membedakan cuma status nikah, pekerjaan, dan kolom kategori.
TETAP_DEWASA = {6: "DKI Jakarta", 7: "Kota Adm. Jakarta Utara", 8: "Koja",
                9: "Lagoa", 13: "Tidak", 14: "Tidak", 15: "Tidak",
                16: "Normal", 17: "Normal", 18: "Normal", 19: "Normal",
                20: "Normal", 21: "Normal", 22: "Normal", 23: "Normal",
                24: "Normal", 27: "Normal", 28: "Normal", 29: "Normal",
                30: "Normal", 31: "Normal", 32: "Normal", 33: "Normal"}
TETAP_REMAJA = TETAP_DEWASA


def _kepala(ws):
    """
    Cari baris kepala, lalu petakan kolomnya dari JUDULNYA.

    Dicari, bukan dihafal: berkas Juni menaruh kepala di baris 4 dan berkas
    Juli di baris 1. Menghafal satu angka berarti berkas berikutnya terbaca
    nol baris tanpa sebab yang jelas.
    """
    for r in range(1, min(12, ws.max_row) + 1):
        judul = {str(ws.cell(row=r, column=k).value or "").strip().upper(): k
                 for k in range(1, ws.max_column + 1)}
        cocok = {}
        for kunci, pola in (("nama", r"NAMA PASIEN"), ("jk", r"^L/P$"),
                            ("lahir", r"TGL"), ("nik", r"^NIK$"),
                            ("alamat", r"^ALAMAT$")):
            for j, k in judul.items():
                if re.search(pola, j):
                    cocok[kunci] = k
                    break
        if len(cocok) == 5:
            return r + 1, cocok
    raise SystemExit("Baris kepala tidak ditemukan - kolom NAMA PASIEN / NIK / TGL tak ada.")


def olah(sumber: Path):
    from openpyxl import load_workbook

    wb = load_workbook(sumber, data_only=True)
    ws = wb.worksheets[0]
    baris_awal, kol = _kepala(ws)
    hari_ini = date.today()

    hasil = {"balita": [], "remaja": [], "dewasa": []}
    catatan: list[str] = []
    terlihat: set[str] = set()
    dibaca = 0

    for r in range(baris_awal, ws.max_row + 1):
        nama_asli = str(ws.cell(row=r, column=kol["nama"]).value or "").strip()
        nik = str(ws.cell(row=r, column=kol["nik"]).value or "").strip()
        if not nama_asli or not nik:
            continue
        dibaca += 1

        if len(nik) != 16 or not nik.isdigit():
            catatan.append(f"baris {r} · {nik} · {nama_asli} · DIBUANG, NIK CACAT: "
                           f"{len(nik)} digit"
                           + ("" if nik.isdigit() else ", ada yang bukan angka"))
            continue
        if nik.endswith("0000"):
            catatan.append(f"baris {r} · {nik} · {nama_asli} · DIBUANG, NIK RUSAK: "
                           "berakhir 0000, digit ke-16 hilang di Excel")
            continue
        if nik in terlihat:
            continue                      # kunjungan kedua orang yang sama
        terlihat.add(nik)

        lahir_berkas = tanggal_kunjungan(ws.cell(row=r, column=kol["lahir"]).value)
        lahir_nik, kelamin_nik = dari_nik(nik)
        if not lahir_nik:
            catatan.append(f"baris {r} · {nik} · {nama_asli} · DIBUANG, nomor ini "
                           "bukan NIK (kemungkinan nomor sementara)")
            continue
        lahir = lahir_berkas or lahir_nik
        if lahir_berkas and lahir_nik != lahir_berkas:
            beda = abs((lahir_berkas - lahir_nik).days) / 365.25
            # Selisih BESAR bukan salah ketik. Salah ketik menggeser hari atau
            # bulan; selisih bertahun-tahun berarti nomornya milik orang lain -
            # dan menuruti NIK di situ memindahkan orang ke kategori yang salah,
            # bahkan mendaftarkannya atas identitas orang lain. Terbukti di
            # berkas Juli: bayi lahir 2026 ber-NIK yang berbunyi 1940.
            if beda >= 2:
                catatan.append(f"baris {r} · {nik} · {nama_asli} · DIBUANG, NIK "
                               f"DIRAGUKAN: berkas {lahir_berkas:%d/%m/%Y} vs NIK "
                               f"{lahir_nik:%d/%m/%Y} - beda {beda:.0f} tahun, "
                               "kemungkinan NIK milik orang lain")
                continue
            catatan.append(f"baris {r} · {nik} · {nama_asli} · tanggal lahir "
                           f"DISESUAIKAN ke NIK: {lahir_berkas:%d/%m/%Y} -> "
                           f"{lahir_nik:%d/%m/%Y}")
            lahir = lahir_nik

        nama = bersihkan_nama(nama_asli)
        if nama != nama_asli:
            catatan.append(f"baris {r} · {nik} · nama dirapikan: {nama_asli!r} -> {nama!r}")

        jk = str(ws.cell(row=r, column=kol["jk"]).value or "").strip().upper()[:1]
        kelamin = kelamin_nik or {"L": "Laki-laki", "P": "Perempuan"}.get(jk, "Laki-laki")
        alamat = " ".join(str(ws.cell(row=r, column=kol["alamat"]).value or "").split()) or "Lagoa"
        umur = umur_tahun(lahir, hari_ini)

        if umur < BALITA_MAKS:
            b = [""] * 18
            b[9], b[12] = alamat, config.NOHP_CADANGAN
            tetap, kunci = TETAP_BALITA, "balita"
        elif umur < REMAJA_MAKS:
            b = [""] * 33
            b[9] = alamat
            b[10] = "Belum Menikah"
            b[11] = "Pelajar"
            # Kolom 25 (kategori) tidak dibaca macro - PARAM_PARAM melompatinya -
            # tapi tetap diisi supaya berkasnya terbaca manusia.
            b[24] = "Remaja"
            # Berat dan tinggi TIDAK dikarang di sini. Tata letak parameter
            # tidak punya kolomnya, dan `Peserta.bb`/`.tb` sudah menurunkan
            # keduanya dari umur dan jenis kelamin.
            tetap, kunci = TETAP_REMAJA, "remaja"
        else:
            b = [""] * 33
            b[9] = alamat
            b[10] = "Menikah" if umur >= 21 else "Belum Menikah"
            b[11] = "Belum/Tidak Bekerja"
            b[24] = "Lansia" if umur >= 60 else "Dewasa"
            tetap, kunci = TETAP_DEWASA, "dewasa"

        b[1], b[2] = nik, nama
        b[3] = lahir.strftime("%d/%m/%Y")
        b[4] = kelamin
        for k, v in tetap.items():
            b[k - 1] = v
        b[0] = str(len(hasil[kunci]) + 1)
        hasil[kunci].append(b)

    catatan.insert(0, f"kepala di baris {baris_awal - 1} · {dibaca} baris kunjungan dibaca · "
                      f"{len(terlihat)} orang berbeda")
    return hasil, catatan


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(2)
    sumber = Path(sys.argv[1])
    hasil, catatan = olah(sumber)

    dasar = re.sub(r"[^a-z0-9]+", "-", sumber.stem.lower()).strip("-")[:40]
    for kunci, baris in hasil.items():
        if not baris:
            print(f"{kunci}: kosong")
            continue
        tujuan = config.DATA / f"datackg-{kunci}-{dasar}.csv"
        with tujuan.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(baris)
        print(f"{kunci:<8}: {len(baris):>4} peserta x {len(baris[0])} kolom -> {tujuan.name}")

    periksa = config.HASIL / "periksa" / f"_pilah-{dasar}.txt"
    periksa.parent.mkdir(parents=True, exist_ok=True)
    periksa.write_text("\n".join(catatan), encoding="utf-8")
    print(f"\n{len(catatan)} catatan -> {periksa}")
    for c in catatan[:6]:
        print(f"   {c}")


if __name__ == "__main__":
    main()
