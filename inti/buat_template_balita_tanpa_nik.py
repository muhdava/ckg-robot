"""
Bangun "Template CKG BALITA TANPA NIK.xlsx" di folder template/.

    python inti\\buat_template_balita_tanpa_nik.py

UNTUK SIAPA
    Bayi yang NIK-nya BELUM TERBIT - di buku kunjungan tertulis "Bayi Nyonya
    ...". Portal mendaftarkannya lewat centang "Tidak punya NIK" di halaman 1,
    lalu tombol "Daftarkan tanpa NIK" di halaman 3.

KENAPA SUSUNANNYA BERBEDA DARI BALITA BIASA
    Tanpa NIK, peserta tidak punya identitas sama sekali, dan portal menutup
    lubang itu dengan MEWAJIBKAN data wali yang lengkap:

        "Pengisian data wali wajib bagi peserta balita dan yang tidak memiliki
         NIK"

    dan di alur ini TIDAK ADA pilihan "Daftarkan tanpa data wali". Jadi kolom
    NIK peserta dibuang, diganti empat kolom wali yang BERURUTAN di kolom 5-8:
    NIK, Nama, Tanggal Lahir, dan Jenis Kelamin - urutan yang sama dengan data
    peserta di sebelah kirinya, sehingga sepasang blok itu mudah dibaca sejajar.

    Nomor HP-nya juga tinggal satu. Di halaman pendaftaran, nomor wali diketik
    ke kolom "No. Whatsapp Aktif" milik PESERTA, lalu nomor wali cukup
    dicentang "Nomor sama dengan peserta" - jadi satu kolom melayani keduanya.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook                                   # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill        # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation    # noqa: E402

import config                                                   # noqa: E402

AKHIR = 2004          # baris terakhir yang ikut diformat & divalidasi

KOLOM = [
    # (judul, lebar, format angka, daftar pilihan, contoh)
    ("No",                     5,  None, "",  "=IF(B5=\"\",\"\",ROW()-4)"),
    ("Nama",                  34,  None, "",  "Bayi Nyonya Contoh"),
    ("Tanggal Lahir",         15,  "@",  "",  "12/01/2024"),
    ("Jenis Kelamin",         14,  None, "Laki-laki,Perempuan", "Laki-laki"),
    ("NIK Wali",              19,  "@",  "",  "3172034105880000"),
    ("Nama Wali",             30,  None, "",  "CONTOH IBU"),
    ("Tanggal Lahir Wali",    18,  "@",  "",  "01/05/1988"),
    ("Jenis Kelamin Wali",    18,  None, "Laki-laki,Perempuan", "Perempuan"),
    ("Provinsi",              14,  None, "",  "DKI Jakarta"),
    ("Kota/Kabupaten",        23,  None, "",  "Kota Adm. Jakarta Utara"),
    ("Kecamatan",             12,  None, "",  "Koja"),
    ("Kelurahan",             12,  None, "",  "Lagoa"),
    ("Alamat",                42,  None, "",  "JL. WALANG TIMUR NO. 30"),
    ("Status Pernikahan",     18,  None, "Belum Menikah", "Belum Menikah"),
    ("Pekerjaan",             16,  None, "Tidak Bekerja,Belum/Tidak Bekerja", "Tidak Bekerja"),
    ("No HP Wali",            16,  "@",  "",  "081210076973"),
    ("Penyandang Disabilitas", 21, None, "Ya,Tidak", "Tidak"),
    ("Berat Lahir",           13,  None, "kurus,normal,gemuk", "normal"),
    ("Berat Saat Ini",        15,  None, "kurus,normal,gemuk", "normal"),
    ("Riwayat Imunisasi",     18,  None, "Ya,Tidak", "Ya"),
    ("Status Gizi",           13,  None, "kurus,normal,gemuk", "normal"),
]

CATATAN = (
    "CKG BALITA TANPA NIK - 21 kolom (A-U). Isi mulai baris 5. EKSPOR: Save As "
    "-> CSV UTF-8, TANPA baris 1-4. PENTING: hapus dulu baris kosong yang hanya "
    "berisi nomor urut. "
    "TIDAK ADA kolom NIK peserta - itu memang yang membedakan template ini; "
    "bayinya didaftarkan lewat centang 'Tidak punya NIK'. Sebagai gantinya "
    "DATA WALI WAJIB LENGKAP (NIK 16 digit, nama, tanggal lahir, jenis kelamin) "
    "karena portal mewajibkannya untuk peserta tanpa NIK. "
    "No HP Wali boleh dikosongkan - bila kosong dipakai nomor hotline "
    "puskesmas. Nomor itu diketik ke kolom WhatsApp peserta di portal, lalu "
    "nomor wali cukup dicentang 'Nomor sama dengan peserta'. "
    "NIK Wali, tanggal, dan No HP diformat sebagai TEKS - jangan diubah ke "
    "Angka, digit ke-16 NIK akan hilang. "
    "Berat & tinggi badan TIDAK diisi di sini: keduanya diturunkan dari tabel "
    "WHO menurut umur dan jenis kelamin."
)

PANDUAN = [
    ("Kolom", "Isi", "Keterangan"),
    ("No", "(otomatis)", "terisi sendiri bila Nama diisi"),
    ("Nama", "teks", "nama bayi, mis. 'Bayi Nyonya Contoh'"),
    ("Tanggal Lahir", "DD/MM/YYYY", "dipakai menghitung umur dan ukuran WHO"),
    ("Jenis Kelamin", "Laki-laki / Perempuan", "milik BAYI"),
    ("NIK Wali", "16 digit", "WAJIB - identitas pengganti karena bayi belum ber-NIK"),
    ("Nama Wali", "teks", "WAJIB"),
    ("Tanggal Lahir Wali", "DD/MM/YYYY", "WAJIB"),
    ("Jenis Kelamin Wali", "Laki-laki / Perempuan", "milik WALI"),
    ("Provinsi..Alamat", "teks alamat domisili", "dipakai penelusuran wilayah bertingkat"),
    ("Status Pernikahan", "Belum Menikah", "milik bayi"),
    ("Pekerjaan", "Tidak Bekerja", "milik bayi"),
    ("No HP Wali", "angka, boleh kosong", "kosong = nomor hotline puskesmas"),
    ("Penyandang Disabilitas", "Ya / Tidak", ""),
    ("Berat Lahir", "kurus/normal/gemuk (PARAM)", "normal = 2600-3500 gram"),
    ("Berat Saat Ini", "kurus/normal/gemuk (PARAM)", "dibandingkan berat ideal WHO"),
    ("Riwayat Imunisasi", "Ya / Tidak", ""),
    ("Status Gizi", "kurus/normal/gemuk (PARAM)", ""),
]


def bangun(tujuan: Path | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    ws["A1"] = "Template Pendaftaran CKG BALITA TANPA NIK (bayi yang NIK-nya belum terbit)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = CATATAN
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    akhir_huruf = chr(ord("A") + len(KOLOM) - 1)
    ws.merge_cells(f"A1:{akhir_huruf}1")
    ws.merge_cells(f"A2:{akhir_huruf}2")
    ws.row_dimensions[2].height = 96

    kepala_isi = PatternFill("solid", fgColor="2D3748")
    for k, (judul, lebar, fmt, pilihan, contoh) in enumerate(KOLOM, 1):
        sel = ws.cell(4, k, judul)
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = kepala_isi
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[chr(ord("A") + k - 1)].width = lebar
        # Baris contoh, supaya bentuk isinya terlihat tanpa membaca panduan.
        ws.cell(5, k, contoh)
        # FORMAT DIKUNCI SAMPAI RIBUAN BARIS, bukan hanya baris contoh: begitu
        # petugas menempel data, kolom NIK yang belum berformat Teks langsung
        # berubah jadi angka dan digit ke-16-nya hilang tanpa suara.
        if fmt:
            for r in range(5, AKHIR + 1):
                ws.cell(r, k).number_format = fmt
        if pilihan:
            dv = DataValidation(type="list", formula1=f'"{pilihan}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{chr(ord('A') + k - 1)}5:{chr(ord('A') + k - 1)}{AKHIR}")
    ws.row_dimensions[4].height = 34
    ws.freeze_panes = "A5"

    # Nomor urut otomatis - mengikuti kolom Nama, sebab kolom NIK tidak ada.
    for r in range(6, AKHIR + 1):
        ws.cell(r, 1, f'=IF(B{r}="","",ROW()-4)')

    w2 = wb.create_sheet("Panduan")
    for i, baris in enumerate(PANDUAN, 1):
        for k, isi in enumerate(baris, 1):
            sel = w2.cell(i, k, isi)
            if i == 1:
                sel.font = Font(bold=True, color="FFFFFF")
                sel.fill = kepala_isi
    for huruf, lebar in (("A", 24), ("B", 32), ("C", 58)):
        w2.column_dimensions[huruf].width = lebar

    # Tujuan boleh ditentukan supaya template bisa dibangun ke tempat lain saat
    # berkas aslinya sedang dibuka Excel - Windows mengunci berkas yang terbuka,
    # dan tanpa jalan keluar ini pembangunannya berhenti dengan PermissionError.
    tujuan = Path(tujuan) if tujuan else (
        config.AKAR / "template" / "Template CKG BALITA TANPA NIK.xlsx")
    tujuan.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tujuan)
    return tujuan


if __name__ == "__main__":
    print(f"Template ditulis: {bangun(Path(sys.argv[1]) if len(sys.argv) > 1 else None)}")
