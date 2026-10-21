"""
Antarmuka web lokal untuk CKG Dewasa/Lansia.

    python server.py          lalu buka http://localhost:8765

Dibuat supaya petugas tidak perlu menyentuh terminal maupun berkas pengaturan:
pilih berkas CSV, ketik rentang, klik jalan, lihat kemajuannya.

BATASAN YANG PERLU DIPAHAMI
  * Browser Playwright terbuka di KOMPUTER TEMPAT SERVER INI BERJALAN, bukan di
    komputer yang membuka halamannya. Login (ada captcha) juga harus dilakukan
    orang yang duduk di depan komputer itu.
  * Karena itu: satu operator, satu komputer. Halaman ini menggantikan terminal,
    bukan menjadikan alatnya bisa dipakai dari jauh.

Memakai pustaka bawaan Python; hanya ekspor Excel yang butuh openpyxl,
dan itu pun diimpor saat dipakai saja - bila tak ada, yang gagal cuma
tombol ekspornya.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import json
import os
import subprocess
import sys
import zipfile
import re
import threading
import webbrowser
import secrets
import socket
import time
from datetime import datetime, timedelta
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import akun
import akun_online
import config
import pembaruan
import pintasan
import konversi
import layanan
import jendela
import paralel
import pekerjaan
import pendaftaran
import peserta
import run as terminal
import sapu_pelayanan
import tatalaksana

# Halaman ada di sebelah kode; data peserta di folder `data/`.
KODE = Path(__file__).resolve().parent
PORTA = 8765

# Daftar pekerja yang sedang berjalan. Satu pekerja = satu browser = satu akun
# portal. Pemakaian satu pekerja hanyalah kasus N=1, jadi tidak ada dua jalur
# kode yang perlahan berbeda perilakunya.
PEKERJA: list = []
_penyiap: threading.Thread | None = None
_kabar = {"pesan": "", "keadaan": "siap"}      # keadaan sebelum pekerja lahir
# Saat pekerjaan dimulai, dalam detik epoch. Dikirim ke halaman supaya
# penghitung waktunya tetap benar walau halamannya dimuat ulang di tengah
# jalan - kalau dihitung di sisi halaman saja, muat ulang membuatnya
# mengulang dari nol dan melaporkan durasi yang lebih pendek dari kenyataan.
_mulai_pada: float = 0.0

# Nama berkas CSV yang sedang/terakhir dijalankan. Dipakai tombol
# "Buat Laporan": laporannya berisi DATA ASLI berkas itu ditambah kolom
# hasil, jadi tanpa mengingat namanya laporan tidak bisa dibuat.
_sumber: str = ""

# ------------------------------------------------------------------- sesi
# Sesi disimpan DI MEMORI saja: begitu aplikasi ditutup, semua sesi hangus.
# Tidak ada berkas sesi yang bisa dicuri, dan tidak ada yang tertinggal di
# komputer bersama.
SESI: dict[str, dict] = {}
UMUR_SESI = timedelta(hours=12)
_gagal: dict[str, list] = {}          # alamat -> [jumlah gagal, waktu terakhir]
BATAS_GAGAL = 5
JEDA_GEMBOK = timedelta(minutes=5)


def _buat_sesi(nama: str, online: dict | None = None) -> str:
    """
    `online` (mode login online): token Supabase + peran. Token Supabase
    hanya disimpan di memori ini, sama seperti sesinya - tidak ke berkas.
    """
    token = secrets.token_urlsafe(32)
    SESI[token] = {"nama": nama, "sampai": datetime.now() + UMUR_SESI, "online": online}
    # Buang sesi kedaluwarsa sekalian, supaya tidak menumpuk.
    for t, d in list(SESI.items()):
        if d["sampai"] < datetime.now():
            del SESI[t]
    return token


_kunci_online = threading.Lock()
# /api/kabar: paling sering sekali per 90 detik ke Supabase, berapa pun tab yang terbuka.
KABAR_JEDA = 90


def _izin_online(sesi: dict) -> dict:
    """
    Tanya ulang Supabase apakah akun ini masih boleh bekerja.

    Dipanggil saat masuk dan setiap kali pekerjaan dimulai, sehingga akun
    yang dinonaktifkan admin berhenti bisa memulai pekerjaan baru paling
    lambat pada tombol "mulai" berikutnya - tanpa menunggu sesi 12 jam habis.
    """
    on = (sesi or {}).get("online")
    if not on:
        return {"ok": True}
    try:
        with _kunci_online:                # penyegaran token tidak boleh balapan
            izin = akun_online.periksa_izin(on["token"])
    except akun_online.GagalOnline as e:
        return {"ok": False, "pesan": e.pesan, "alasan": e.alasan}
    if izin.get("ok"):
        on["peran"] = izin.get("peran", on.get("peran"))
        on["pengumuman"] = izin.get("pengumuman", "")
    if "versi_terbaru" in izin:
        on["pembaruan"] = {"versi": izin.get("versi_terbaru") or "", "sha256": izin.get("sha256_terbaru") or "",
                           "ukuran": izin.get("ukuran_terbaru") or 0,
                           "catatan": izin.get("catatan_terbaru") or ""}
    on["wajib_perbarui"] = izin.get("alasan") == "versi_lama"
    return izin


def _pembaruan_tersedia(sesi: dict | None) -> dict | None:
    on = (sesi or {}).get("online") or {}
    p = on.get("pembaruan") or {}
    if not pembaruan.lebih_baru(p.get("versi", ""), config.VERSI):
        return None
    return {"versi": p["versi"], "catatan": p.get("catatan", ""), "wajib": bool(on.get("wajib_perbarui")),
            "versi_ini": config.VERSI}


# ------------------------------------------------------------ pembaruan
# Satu pembaruan pada satu waktu. Keadaannya dibaca halaman lewat
# /api/perbarui-status; sesudah 'mulai_ulang' server lama berhenti melayani
# dan halaman menunggu server baru menjawab.
_pembaruan_keadaan: dict = {"tahap": "", "pesan": "", "persen": 0, "galat": ""}
_kunci_pembaruan = threading.Lock()
_mulai_ulang = False
_SERVER: list = []


def _lapor_pembaruan(tahap: str, pesan: str, persen: int) -> None:
    _pembaruan_keadaan.update(tahap=tahap, pesan=pesan, persen=persen)


def _kerjakan_pembaruan(sesi: dict) -> None:
    on = sesi["online"]
    try:
        with _kunci_online:
            akun_online.akses_segar(on["token"])
        hasil = pembaruan.jalankan(on["token"], on["pembaruan"], _lapor_pembaruan)
    except (pembaruan.GagalPembaruan, akun_online.GagalOnline) as e:
        _pembaruan_keadaan.update(tahap="gagal", galat=getattr(e, "pesan", None) or str(e), persen=0)
        return
    except Exception as e:                              # jangan sampai utasnya mati diam-diam
        _catat_galat("pembaruan", e)
        _pembaruan_keadaan.update(tahap="gagal", galat=f"{type(e).__name__}: {e}", persen=0)
        return
    catatan = f"Versi {hasil['versi']} terpasang ({hasil['berkas']} berkas)."
    if hasil.get("galat_pustaka"):
        catatan += " Pemasangan pustaka gagal - jalankan pasang.bat sekali."
    _pembaruan_keadaan.update(tahap="mulai_ulang", pesan=catatan + " Memulai ulang aplikasi…", persen=100,
                              versi=hasil["versi"])
    # Beri waktu halaman membaca keadaan 'mulai_ulang' sebelum server berhenti.
    threading.Timer(1.5, _hentikan_untuk_mulai_ulang).start()


def _hentikan_untuk_mulai_ulang() -> None:
    """Hentikan server ini; main() lalu menjalankan server versi baru (lihat main)."""
    global _mulai_ulang
    _mulai_ulang = True
    for s in _SERVER:
        try:
            s.shutdown()
        except Exception:
            pass


# ------------------------------------------------------------ pemakaian
# JUMLAH peserta per kategori dikirim ke Supabase (login online), supaya admin
# melihat siapa yang aktif dan - lebih penting - melihat lonjakan "gagal" di
# semua petugas sekaligus: tanda portal Kemenkes berubah. Hanya ANGKA; tidak
# ada NIK, nama, atau data peserta. Yang dikirim SELISIH sejak kiriman terakhir,
# jadi hasil putaran ulang (gagal -> tuntas) ikut terkoreksi.
KATEGORI_PEMAKAIAN = {
    "tuntas": "tuntas", "layanan_selesai": "tuntas",
    "sudah_pelayanan": "beres_lain", "ckg_umum": "beres_lain", "ckg_sekolah": "beres_lain",
    "puskesmas_berbeda": "beres_lain", "sudah_tatalaksana": "beres_lain",
    "tidak_valid": "data_salah", "tidak_ditemukan": "data_salah", "sekolah_berbeda": "data_salah",
    "kelas_berbeda": "data_salah",
}
_pemakaian_terkirim: dict[int, dict] = {}      # id(kerja) -> hitungan yang sudah terkirim
_sesi_kerja: dict | None = None                 # sesi yang memulai pekerjaan terakhir
_kunci_pemakaian = threading.Lock()


def _hitung(kerja) -> dict:
    h = {"tuntas": 0, "beres_lain": 0, "data_salah": 0, "gagal": 0}
    for d in list(kerja.hasil):
        s = (d.get("status") or "").strip()
        if s:
            h[KATEGORI_PEMAKAIAN.get(s, "gagal")] += 1
    return h


def _kirim_pemakaian() -> None:
    """Kirim selisih hitungan semua pekerja. Gagal kirim = dicoba lagi lain kali."""
    on = (_sesi_kerja or {}).get("online")
    if not on or not akun_online.aktif():
        return
    with _kunci_pemakaian:
        selisih = {"tuntas": 0, "beres_lain": 0, "data_salah": 0, "gagal": 0}
        sekarang = {}
        for p in list(PEKERJA):
            h = _hitung(p.kerja)
            sekarang[id(p.kerja)] = h
            lama = _pemakaian_terkirim.get(id(p.kerja), {})
            for k in selisih:
                selisih[k] += h[k] - lama.get(k, 0)
        if not any(selisih.values()):
            return
        try:
            with _kunci_online:
                r = akun_online.rpc("catat_pemakaian", {
                    "p_tanggal": datetime.now().strftime("%Y-%m-%d"), "p_tuntas": selisih["tuntas"],
                    "p_beres_lain": selisih["beres_lain"], "p_data_salah": selisih["data_salah"],
                    "p_gagal": selisih["gagal"], "p_versi": config.VERSI}, on["token"])
        except akun_online.GagalOnline:
            return
        if r.get("ok"):
            _pemakaian_terkirim.update(sekarang)


def _pengirim_pemakaian() -> None:
    while True:
        time.sleep(60)
        try:
            _kirim_pemakaian()
        except Exception as e:                  # jangan sampai utasnya mati diam-diam
            _catat_galat("kirim pemakaian", e)


def _paket_lokal() -> list[dict]:
    """ZIP buatan buat_paket.bat di folder induk, terbaru dulu."""
    hasil = []
    for f in sorted(config.PROYEK.glob("CKG-Robot-*.zip"), key=lambda f: f.stat().st_mtime, reverse=True)[:8]:
        try:
            with zipfile.ZipFile(f) as zf:
                versi = pembaruan.versi_dalam_zip(zf)
        except Exception:
            versi = ""
        hasil.append({"nama": f.name, "versi": versi, "ukuran": f.stat().st_size,
                      "waktu": datetime.fromtimestamp(f.stat().st_mtime).strftime("%d-%m-%Y %H:%M")})
    return hasil


def _rpc_admin(sesi: dict, fungsi: str, argumen: dict) -> dict:
    on = (sesi or {}).get("online") or {}
    if on.get("peran") != "admin":
        return {"ok": False, "pesan": "Hanya admin."}
    try:
        with _kunci_online:
            return akun_online.rpc(fungsi, argumen, on["token"])
    except akun_online.GagalOnline as e:
        return {"ok": False, "pesan": e.pesan}


# ------------------------------------------------------------------ ekspor
# Warna per status. Dikelompokkan menurut ARTINYA, bukan menurut namanya:
# hijau = beres, kelabu = memang sudah dilayani, kuning = urusan DATA yang
# perlu diperiksa manusia, merah = pekerjaan yang gagal dan boleh diulang.
WARNA_STATUS = {
    "tuntas":           ("C6F6D5", "1C4532"),
    "layanan_selesai":  ("C6F6D5", "1C4532"),
    "sudah_pelayanan":  ("E2E8F0", "2D3748"),
    # Sudah diperiksa lewat CKG Umum - selesai, hanya bukan lewat
    # jalur sekolah. Biru muda supaya bisa dipilah sekilas.
    "ckg_umum":         ("BEE3F8", "2A4365"),
    # Sudah selesai diperiksa, hanya di sekolah/kelas yang BUKAN seperti di
    # CSV. Bukan kegagalan alat dan bukan pula "beres di tempatnya" - ungu
    # supaya langsung terlihat sebagai urusan data yang perlu dibetulkan.
    "sekolah_berbeda":  ("E9D8FD", "44337A"),
    "kelas_berbeda":    ("E9D8FD", "44337A"),
    # Sudah diperiksa lewat jalur CKG Sekolah - selesai, hanya bukan lewat
    # pintu CKG Umum. Hijau kebiruan supaya terpisah dari ckg_umum.
    "ckg_sekolah":      ("B2F5EA", "234E52"),
    # Diperiksa puskesmas lain - catatannya bukan milik kita.
    "puskesmas_berbeda": ("E9D8FD", "44337A"),
    # Terdaftar tapi belum tentu hadir - bukan kegagalan, jadi bukan merah.
    "sudah_terdaftar":  ("E2E8F0", "2D3748"),
    "tidak_valid":      ("FEEBC8", "7B341E"),
    "tidak_ditemukan":  ("FEF5C7", "744210"),
    # Portal tidak memuat daftar pilihannya - masalah di sana, bukan di
    # sini, dan bisa berhasil kalau diulang nanti. Kuning peringatan.
    "daftar_kosong":    ("FEF5C7", "744210"),
    "uji":              ("BEE3F8", "2A4365"),
    "layanan_sebagian": ("FED7AA", "7B341E"),
    # Tatalaksana yang sudah dicatat sebelumnya (petugas atau jalan lain).
    "sudah_tatalaksana": ("E2E8F0", "2D3748"),
}
WARNA_GAGAL = ("FED7D7", "742A2A")

# Kalimat pendek untuk kolom "Hasil" di berkas laporan. Kata statusnya sendiri
# tetap ada di kolom "Status" - itu yang enak disaring; yang ini yang enak
# dibaca petugas sekolah atau pengelola program, yang tidak tahu-menahu soal
# nama status di dalam macro.
ARTI_STATUS = {
    "tuntas":            "Tuntas - seluruh form terisi",
    "layanan_selesai":   "Seluruh layanan sudah ditutup",
    "sudah_pelayanan":   "Sudah dilayani",
    "ckg_umum":          "Sudah diperiksa lewat CKG Umum",
    "ckg_sekolah":       "Sudah diperiksa lewat CKG Sekolah",
    "puskesmas_berbeda": "Sudah diperiksa PUSKESMAS LAIN",
    "sekolah_berbeda":   "Sudah diperiksa di SEKOLAH LAIN",
    "kelas_berbeda":     "Sudah diperiksa, KELAS berbeda",
    "sudah_terdaftar":   "Sudah terdaftar, belum dilayani",
    "tidak_valid":       "Ditolak portal - NIK/data perlu dibetulkan",
    "tidak_ditemukan":   "Tidak ada di daftar portal",
    "daftar_kosong":     "Portal tidak memuat daftarnya",
    "layanan_sebagian":  "Diperiksa tapi form belum lengkap",
    "sudah_tatalaksana": "Tatalaksana sudah dicatat sebelumnya",
    "gagal_isi":         "Form tatalaksana tidak dikirim (isian tidak cocok)",
    "":                  "Belum dijalankan",
}


def arti_status(status: str) -> str:
    """Kalimat untuk satu status. Yang belum dikenal dilaporkan apa adanya."""
    if status in ARTI_STATUS:
        return ARTI_STATUS[status]
    return f"Gagal di tengah jalan ({status}) - bisa diulang"          # gagal_*, error, tanpa_tiket, dll.


# Status yang berarti murid SUDAH diperiksa (di sini atau di tempat lain).
BERES = {"tuntas", "layanan_selesai", "sudah_pelayanan", "ckg_umum", "ckg_sekolah", "puskesmas_berbeda"}
HB_BATAS = 12.0          # WHO: anemia remaja 12-14 th bila Hb < 12 g/dL


def langkah_tindakan(status: str) -> str:
    if status == "tidak_valid":
        return ("Cocokkan NIK, nama, dan tanggal lahir dengan Kartu Keluarga. Betulkan di berkas, lalu "
                "jalankan ulang (pengaturan 'ulang tidak valid = ya').")
    if status == "tidak_ditemukan":
        return "Periksa NIK, nama sekolah, dan kelas - harus sama persis dengan tulisan di portal."
    if status in ("sekolah_berbeda", "kelas_berbeda"):
        return "Murid tercatat di sekolah/kelas lain di portal. Betulkan datanya, atau abaikan bila memang pindah."
    if status in ("", "belum dijalankan"):
        return "Belum dijalankan - jalankan nomor ini."
    return "Gagal sementara (portal lambat/macet) - jalankan ulang nomor ini."


def _lembar_sekolah(wb, kepala: list[str], semua: list[list], per_baris: dict) -> None:
    """Lembar tambahan untuk berkas CKG Sekolah. Lihat buat_laporan_xlsx."""
    from openpyxl.styles import Font, PatternFill

    kol = {n: kepala.index(n) for n in ("NIK", "Nama", "Jenis Kelamin", "Nama Sekolah", "Kelas", "Hemoglobin")
           if n in kepala}
    gelap = PatternFill("solid", fgColor="2D3748")
    putih = Font(bold=True, color="FFFFFF")

    def sel(b, nama):
        k = kol.get(nama)
        return (b[k] if k is not None and k < len(b) else "") or ""

    murid = []
    for i, b in enumerate(semua):
        d = per_baris.get(i + 1, {})
        status = (d.get("status") or "").strip()
        murid.append({"no": i + 1, "b": b, "status": status, "ket": (d.get("keterangan") or "").strip()})
    dijalankan = [m for m in murid if m["status"]]
    urut_kelas = lambda k: (int("".join(ch for ch in k if ch.isdigit()) or 0), k)  # noqa: E731
    kelas = sorted({sel(m["b"], "Kelas") or "(tanpa kelas)" for m in murid}, key=urut_kelas)

    # ---------------------------------------------------- Rekap per Kelas
    ws = wb.create_sheet("Rekap per Kelas", 0)
    sekolah = sorted({sel(m["b"], "Nama Sekolah") for m in murid if sel(m["b"], "Nama Sekolah")})
    ws["A1"] = "Rekap hasil CKG Sekolah" + (f" - {', '.join(sekolah)}" if sekolah else "")
    ws["A1"].font = Font(bold=True, size=13)
    beres = sum(1 for m in dijalankan if m["status"] in BERES)
    ws["A2"] = (f"{len(murid)} murid di berkas, {len(dijalankan)} sudah dijalankan. Sudah beres (diperiksa di "
                f"sini atau di tempat lain): {beres} murid"
                + (f" ({round(100 * beres / len(dijalankan), 1)}% dari yang dijalankan)." if dijalankan else "."))
    ws.append([])
    ws.append(["Status", "Artinya"] + kelas + ["Jumlah"])
    for k in range(1, len(kelas) + 4):
        ws.cell(4, k).font, ws.cell(4, k).fill = putih, gelap
    urutan = ["tuntas", "sudah_pelayanan", "ckg_umum", "ckg_sekolah", "puskesmas_berbeda", "sekolah_berbeda",
              "kelas_berbeda", "tidak_valid", "tidak_ditemukan"]
    semua_status = sorted({m["status"] or "belum dijalankan" for m in murid},
                          key=lambda s: (urutan.index(s) if s in urutan else 50 if s != "belum dijalankan" else 99, s))
    for s in semua_status:
        cocok = [m for m in murid if (m["status"] or "belum dijalankan") == s]
        baris = [s, arti_status("" if s == "belum dijalankan" else s)]
        baris += [sum(1 for m in cocok if (sel(m["b"], "Kelas") or "(tanpa kelas)") == k) for k in kelas]
        ws.append(baris + [len(cocok)])
        latar, huruf = WARNA_STATUS.get(s, WARNA_GAGAL if s != "belum dijalankan" else ("FFFFFF", "4A5568"))
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=latar)
        ws.cell(ws.max_row, 1).font = Font(bold=True, color=huruf)
    ws.append(["JUMLAH", ""] + [sum(1 for m in murid if (sel(m["b"], "Kelas") or "(tanpa kelas)") == k)
                                for k in kelas] + [len(murid)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 46

    # ---------------------------------------------------- Perlu Tindakan
    pt = wb.create_sheet("Perlu Tindakan", 1)
    pt.append(["Baris", "Kelas", "Nama", "NIK", "Status", "Langkah", "Keterangan robot (ringkas)"])
    for c in pt[1]:
        c.font, c.fill = putih, gelap
    for m in dijalankan:
        if m["status"] in BERES:
            continue
        pt.append([m["no"], sel(m["b"], "Kelas"), sel(m["b"], "Nama"), sel(m["b"], "NIK"), m["status"],
                   langkah_tindakan(m["status"]), m["ket"][-250:]])
        pt.cell(pt.max_row, 4).number_format = "@"
    if pt.max_row == 1:
        pt.append(["", "", "Tidak ada - semua murid yang dijalankan sudah beres."])
    for k, w in zip("ABCDEFG", (7, 10, 30, 19, 18, 60, 80)):
        pt.column_dimensions[k].width = w

    # ---------------------------------------------------- Hb Rendah
    if "Hemoglobin" in kol:
        hb = wb.create_sheet("Hb Rendah", 2)
        hb.append(["Baris", "Kelas", "Nama", "NIK", "Jenis Kelamin", "Hb (g/dL)", "Derajat (WHO)"])
        for c in hb[1]:
            c.font, c.fill = putih, gelap
        rendah = []
        for m in murid:
            try:
                nilai = float(str(sel(m["b"], "Hemoglobin")).replace(",", "."))
            except ValueError:
                continue
            if 0 < nilai < HB_BATAS:
                rendah.append((nilai, m))
        for nilai, m in sorted(rendah, key=lambda x: x[0]):
            derajat = ("anemia berat (< 8)" if nilai < 8 else "anemia sedang (8 - 10,9)" if nilai < 11
                       else "anemia ringan (11 - 11,9)")
            hb.append([m["no"], sel(m["b"], "Kelas"), sel(m["b"], "Nama"), sel(m["b"], "NIK"),
                       sel(m["b"], "Jenis Kelamin"), nilai, derajat])
            hb.cell(hb.max_row, 4).number_format = "@"
        if not rendah:
            hb.append(["", "", f"Tidak ada murid dengan Hb di bawah {HB_BATAS:g} g/dL."])
        hb.append([])
        hb.append(["", "", "Catatan: Hb yang kosong di berkas diisi 13 oleh konverter, jadi tidak muncul di sini."])
        for k, w in zip("ABCDEFG", (7, 10, 30, 19, 13, 10, 26)):
            hb.column_dimensions[k].width = w
    wb.active = 0


def buat_laporan_xlsx(nama_csv: str, hasil: list[dict]) -> tuple[bytes, str]:
    """
    Susun LAPORAN: berkas yang dijalankan, ditambah satu kolom "Hasil".

    Bedanya dengan "Unduh Excel": yang itu berisi hasil saja - tujuh kolom
    ringkas untuk dikirim saat minta bantuan. Yang ini berisi DATA ASLINYA
    UTUH, persis kolom yang dikirim ke portal, dengan kolom "Hasil" disisipkan
    tepat di kanan "Jenis Kelamin". Itulah bentuk yang dipakai melapor balik ke
    sekolah atau pengelola program: mereka mengenali barisnya sendiri.

    Baris hasil disambungkan lewat NOMOR BARIS CSV, bukan NIK. Nomor baris
    tidak pernah kembar, sedangkan NIK bisa - dan pada sasaran tanpa-NIK
    kolomnya bahkan tidak ada.

    Susunan kolomnya dikenali dari LEBARNYA (18/21/24/33/34) lalu dicocokkan
    ke katalog `ambil_data_umum.TATA_LETAK`. Berkas dengan lebar tak dikenal
    tetap bisa dilaporkan - judulnya saja yang jadi "Kolom 1..N".
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    import ambil_data_umum as umum

    berkas = config.DATA / nama_csv
    with open(berkas, encoding="utf-8-sig", newline="") as fh:
        semua = [b for b in csv.reader(fh) if any((x or "").strip() for x in b)]
    if not semua:
        raise ValueError("berkas sumber kosong")

    lebar = max(len(b) for b in semua)
    kepala = next(([k[0] for k in v] for v in umum.TATA_LETAK.values() if len(v) == lebar),
                  [f"Kolom {i}" for i in range(1, lebar + 1)])
    k_jk = (kepala.index("Jenis Kelamin") + 1) if "Jenis Kelamin" in kepala else len(kepala)
    judul = kepala[:k_jk] + ["Hasil"] + kepala[k_jk:] + ["Status", "Keterangan", "Tiket", "Waktu"]

    per_baris = {}
    for d in hasil:
        try:
            per_baris[int(d.get("baris") or 0)] = d
        except (TypeError, ValueError):
            continue

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    isi_kepala = PatternFill("solid", fgColor="2D3748")
    for k, t in enumerate(judul, 1):
        sel = ws.cell(4, k, t)
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = isi_kepala
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 34
    ws.freeze_panes = f"{get_column_letter(k_jk + 2)}5"

    n_ada = 0
    for i, b in enumerate(semua):
        b = list(b) + [""] * (lebar - len(b))
        d = per_baris.get(i + 1, {})          # baris_csv dihitung mulai 1
        status = (d.get("status") or "").strip()
        n_ada += bool(status)
        latar, huruf = WARNA_STATUS.get(status, WARNA_GAGAL if status else ("FFFFFF", "4A5568"))
        nilai = (b[:k_jk] + [arti_status(status)] + b[k_jk:]
                 + [status or "belum dijalankan", (d.get("keterangan") or "").strip(),
                    (d.get("tiket") or "").strip(), (d.get("waktu") or "").strip()])
        for k, v in enumerate(nilai, 1):
            ws.cell(5 + i, k, v)
        sel = ws.cell(5 + i, k_jk + 1)
        sel.fill = PatternFill("solid", fgColor=latar)
        sel.font = Font(bold=True, color=huruf)
    # NIK, tanggal, dan nomor HP tetap TEKS: dibaca Excel sebagai angka, digit
    # terakhir NIK hilang dan angka nol di depan nomor HP ikut lenyap.
    for nama_kol in ("NIK", "NIK Wali", "Tanggal Lahir", "Tanggal Lahir Wali",
                     "No HP", "No HP Wali"):
        if nama_kol in judul:
            k = judul.index(nama_kol) + 1
            for r in range(5, len(semua) + 5):
                ws.cell(r, k).number_format = "@"
    for k, lebar_kol in enumerate([5, 19, 30, 14, 13, 34], 1):
        ws.column_dimensions[get_column_letter(k)].width = lebar_kol
    ws.column_dimensions[get_column_letter(len(judul) - 2)].width = 70
    ws.auto_filter.ref = f"A4:{get_column_letter(len(judul))}{len(semua) + 4}"

    ws["A1"] = f"LAPORAN HASIL - {nama_csv}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"{len(semua)} baris data, {n_ada} sudah punya hasil. Kolom Hasil ada tepat "
                "di kanan Jenis Kelamin; kolom lainnya sama persis dengan berkas yang "
                "dijalankan. Baris tanpa hasil berarti belum dijalankan pada rentang ini. "
                "Berkas ini untuk DIBACA - kolom tambahannya membuat susunannya tidak lagi "
                "cocok untuk dikonversi.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"A1:{get_column_letter(len(judul))}1")
    ws.merge_cells(f"A2:{get_column_letter(len(judul))}2")
    ws.row_dimensions[2].height = 46

    # Berkas CKG SEKOLAH: lembar rekap untuk sekolah & puskesmas, supaya
    # laporan seperti SMPN 136 tidak perlu disusun manual.
    if "Kelas" in kepala and "Nama Sekolah" in kepala:
        _lembar_sekolah(wb, kepala, [list(b) + [""] * (lebar - len(b)) for b in semua], per_baris)

    keluar = BytesIO()
    wb.save(keluar)
    return keluar.getvalue(), Path(nama_csv).stem


def buat_xlsx(baris: list[dict]) -> bytes:
    """
    Susun berkas Excel berisi hasil satu pekerjaan.

    openpyxl diimpor DI SINI, bukan di kepala berkas: kalau pustakanya belum
    terpasang, yang gagal hanya tombol ekspor - bukan seluruh aplikasi.
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Hasil"
    # KETERANGAN IKUT DIEKSPOR. Tanpa kolom ini berkas hasil tidak bisa
    # menjelaskan apa pun: "layanan_sebagian" dan "gagal_detail" terlihat sama
    # saja entah karena portal menggantung, form yang tak ter-render, atau
    # percobaan ulang yang sudah habis - padahal ketiganya butuh tindakan
    # berbeda. Jejaknya sudah lama ada di CSV; hanya ekspor Excel yang
    # membuangnya, dan itulah berkas yang dikirim orang saat minta bantuan.
    judul = ["No", "NIK", "Nama", "Tanggal Lahir", "Status", "Tiket", "Keterangan"]
    ws.append(judul)

    kepala_isi = PatternFill("solid", fgColor="1F7A5C")
    kepala_huruf = Font(bold=True, color="FFFFFF")
    garis = Side(style="thin", color="D9DEE3")
    kotak = Border(left=garis, right=garis, top=garis, bottom=garis)
    for i in range(1, len(judul) + 1):
        sel = ws.cell(row=1, column=i)
        sel.fill, sel.font, sel.border = kepala_isi, kepala_huruf, kotak
        sel.alignment = Alignment(horizontal="center", vertical="center")

    for nomor, b in enumerate(baris, start=1):
        status = (b.get("status") or "").strip()
        ws.append([nomor,
                   str(b.get("nik") or ""),
                   b.get("nama") or "",
                   b.get("tgl_lahir") or "",
                   status,
                   str(b.get("tiket") or ""),
                   (b.get("keterangan") or "").strip(" |")])
        r = ws.max_row
        for i in range(1, len(judul) + 1):
            ws.cell(row=r, column=i).border = kotak
        # NIK sebagai teks: 16 digit melampaui ketelitian angka Excel, dan
        # kalau dianggap angka digit terakhirnya berubah jadi nol.
        ws.cell(row=r, column=2).number_format = "@"
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")

        isi, huruf = WARNA_STATUS.get(status, WARNA_GAGAL if status else ("FFFFFF", "1A1D20"))
        sel = ws.cell(row=r, column=5)
        sel.fill = PatternFill("solid", fgColor=isi)
        sel.font = Font(bold=True, color=huruf)
        sel.alignment = Alignment(horizontal="center")

    for kolom, lebar in zip("ABCDEFG", (6, 22, 34, 15, 20, 14, 110)):
        ws.column_dimensions[kolom].width = lebar
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(judul))}{ws.max_row}"

    # --- Rekap di LEMBAR TERSENDIRI ---
    # Bukan blok di atas tabel: menyisipkan baris di atas judul akan merusak
    # baris beku dan filter otomatis, yang justru paling sering dipakai saat
    # memeriksa hasil.
    jumlah: dict[str, int] = {}
    for b in baris:
        s = (b.get("status") or "").strip() or "(tanpa status)"
        jumlah[s] = jumlah.get(s, 0) + 1

    # Urutan tetap: yang beres dulu, lalu yang perlu dicek, lalu yang gagal.
    urutan_tetap = ["tuntas", "layanan_selesai", "sudah_pelayanan",
                    "ckg_umum", "ckg_sekolah", "puskesmas_berbeda",
                    "sekolah_berbeda", "kelas_berbeda",
                    "sudah_terdaftar", "tidak_valid",
                    "tidak_ditemukan", "daftar_kosong",
                    "layanan_sebagian", "uji"]
    def kunci(s: str):
        i = urutan_tetap.index(s) if s in urutan_tetap else 99
        return (i, -jumlah[s], s)

    rk = wb.create_sheet("Rekap")
    rk.append(["Status", "Jumlah"])
    for i in (1, 2):
        sel = rk.cell(row=1, column=i)
        sel.fill, sel.font, sel.border = kepala_isi, kepala_huruf, kotak
        sel.alignment = Alignment(horizontal="center", vertical="center")

    for status in sorted(jumlah, key=kunci):
        rk.append([status, jumlah[status]])
        r = rk.max_row
        isi, huruf = WARNA_STATUS.get(status, WARNA_GAGAL)
        for i in (1, 2):
            sel = rk.cell(row=r, column=i)
            sel.border = kotak
            sel.fill = PatternFill("solid", fgColor=isi)
            sel.font = Font(bold=(i == 2), color=huruf)
        rk.cell(row=r, column=2).alignment = Alignment(horizontal="center")

    rk.append(["Total peserta", len(baris)])
    r = rk.max_row
    for i in (1, 2):
        sel = rk.cell(row=r, column=i)
        sel.border, sel.font = kotak, Font(bold=True)
        sel.fill = PatternFill("solid", fgColor="EDF2F7")
    rk.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    rk.column_dimensions["A"].width = 24
    rk.column_dimensions["B"].width = 10
    rk.freeze_panes = "A2"

    tampung = BytesIO()
    wb.save(tampung)
    return tampung.getvalue()


# ---------------------------------------------------------------- pekerjaan
def daftar_csv() -> list[dict]:
    """Berkas CSV di folder ini yang tampak seperti data peserta."""
    keluar = []
    for f in sorted(config.DATA.glob("*.csv")):
        if f.name.startswith("_"):        # berkas '-ditolak' / '-perlu-dicek'
            continue
        try:
            jumlah = len(peserta.baca_csv(f))
        except Exception:
            continue
        if jumlah:
            keluar.append({"nama": f.name, "jumlah": jumlah})
    return keluar


DIIZINKAN = (".xlsx", ".xlsm", ".csv")


def _laporan_gagal(pesan: str, nama: str = "") -> dict:
    """
    Laporan konversi yang berisi satu peringatan saja.

    NAMANYA JANGAN `_gagal`: nama itu sudah dipakai kamus pembatas percobaan
    masuk di atas, dan menimpanya membuat `_gagal.get(alamat)` melempar
    AttributeError - siapa pun jadi tidak bisa login sama sekali.
    """
    return {"peringatan": [pesan], "catatan": [], "jumlah": 0,
            "berkas": nama, "sebutan": "", "alasan": "", "tujuan": ""}


def _simpan_unggahan(nama: str, isi_b64: str) -> tuple[Path | None, str]:
    """
    Simpan berkas yang dipilih lewat halaman ke folder `mentah/`.

    Disimpan, bukan diolah langsung dari memori: pengekstraknya membaca lewat
    openpyxl yang menerima jalur, dan salinannya juga berguna - bila hasilnya
    meleset, berkas persis yang dipakai masih ada untuk diperiksa.

    NAMANYA DIPANGKAS jadi nama berkas saja. Nama yang dikirim halaman tidak
    boleh dipercaya sebagai jalur: "../../inti/server.py" akan menimpa kode
    aplikasi sendiri.
    """
    aman = Path(nama).name.strip()
    if not aman:
        return None, "Nama berkas kosong."
    if Path(aman).suffix.lower() not in DIIZINKAN:
        return None, (f"Jenis berkas '{Path(aman).suffix}' tidak didukung. "
                      "Yang bisa: " + ", ".join(DIIZINKAN))
    try:
        isi = base64.b64decode(isi_b64, validate=True)
    except Exception:
        return None, "Isi berkas tidak bisa dibaca (base64 rusak)."
    config.siapkan_folder()
    tujuan = config.MENTAH / aman
    tujuan.write_bytes(isi)
    return tujuan, ""


def _konversi(data: dict) -> dict:
    """Jalankan konversi satu berkas mentah lalu kembalikan laporannya."""
    nama = (data.get("berkas") or "").strip()
    if data.get("isi"):
        # Berkas dipilih langsung dari komputer - tidak perlu dipindah dulu.
        berkas, salah = _simpan_unggahan(nama, data["isi"])
        if salah:
            return _laporan_gagal("BERHENTI: " + salah, nama)
    else:
        # Nama berkas datang dari daftar yang KAMI susun, tapi tetap dibatasi ke
        # folder `mentah/`: menerima jalur apa adanya berarti halaman ini bisa
        # dipakai membaca berkas mana pun di komputer.
        berkas = (config.MENTAH / Path(nama).name)
        if not nama or not berkas.exists():
            return _laporan_gagal(
                f"BERHENTI: berkas '{nama}' tidak ada di folder mentah/.", nama)
    kelas = [int(x) for x in re.findall(r"\d+", data.get("kelas") or "")]
    return konversi.jalankan(berkas, kelas=kelas,
                             sekolah=(data.get("sekolah") or "").strip(),
                             bentuk=(data.get("bentuk") or "auto").strip(),
                             ikut_nik=bool(data.get("ikut_nik", True)),
                             tahan_kosong=not data.get("ikut_kosong", False))


def _siapkan_dan_jalankan(nama_csv: str, rentang: str, jumlah_pekerja: int) -> None:
    """
    Baca berkas, bagi datanya, lalu lepas para pekerja.

    Pembagian dilakukan SEKALI di sini, bukan oleh masing-masing pekerja:
    peserta yang sama dikerjakan dua pekerja berarti pendaftaran ganda di
    portal, dan itu tidak bisa dibatalkan.
    """
    global PEKERJA, _mulai_pada, _sumber
    _mulai_pada = time.time()
    _sumber = nama_csv
    try:
        berkas = config.DATA / nama_csv
        semua = peserta.baca_csv(berkas)
        diminta = terminal.urai_nomor(rentang)
        batch = [p for p in semua if p.baris_csv in set(diminta)] if diminta else []
        if not batch:
            _kabar.update(keadaan="gagal", pesan="Tidak ada baris yang cocok dengan rentang itu.")
            return

        batch, dilewati, ditolak = terminal.pilah_yang_dilewati(batch)
        if not batch:
            _kabar.update(keadaan="selesai",
                          pesan="Semua peserta pada rentang itu sudah selesai atau "
                                "sudah ditolak portal sebelumnya.")
            return

        PEKERJA = paralel.siapkan(batch, jumlah_pekerja)
        alur = {"pelayanan": "Pelayanan → Selesai",
                "daftar_konfirmasi": "Pendaftaran → Konfirmasi Hadir → Selesai",
                "pelayanan_sapu": "Pelayanan → Selesai (sapu tab Selesai)",
                "daftar_pelayanan": "Pendaftaran → Pelayanan → Selesai"}.get(
                    config.TAHAP, "Pendaftaran → Selesai")
        for p in PEKERJA:
            p.kerja.lapor(f"Berkas {berkas.name} · alur: {alur} · jenis: {config.KATEGORI}")
            if dilewati:
                p.kerja.lapor(f"{dilewati} peserta dilewati — sudah selesai sebelumnya.")
            if ditolak:
                p.kerja.lapor(f"{ditolak} peserta dilewati — sebelumnya ditolak portal "
                              "(tidak valid). Betulkan datanya, lalu nyalakan "
                              "'ulang tidak valid' di pengaturan.")
            p.mulai()
        _kabar.update(keadaan="berjalan", pesan="")
    except Exception as e:
        _kabar.update(keadaan="gagal",
                      pesan=f"{type(e).__name__}: {' '.join(str(e).split())[:250]}")


def _mulai_sapu(data: dict) -> dict:
    """
    Alur "Sapu Belum Pemeriksaan": pelayanan dewasa/lansia tanpa berkas.
    Satu pekerja - semua browser melihat daftar yang sama (lihat _mulai_tatalaksana).
    """
    global PEKERJA, _mulai_pada, _sumber
    jawaban, galat = sapu_pelayanan.urai_jawaban(data.get("sapu"))
    if galat:
        return {"ok": False, "pesan": galat}
    try:
        jumlah = int(data.get("jumlah"))
    except (TypeError, ValueError):
        return {"ok": False, "pesan": "Isi jumlah peserta (angka)."}
    if not 1 <= jumlah <= config.SAPU_MAKS:
        return {"ok": False, "pesan": f"Jumlah peserta 1 sampai {config.SAPU_MAKS} per jalan."}
    config.TAHAP = "pelayanan"
    config.KATEGORI = "dewasa"
    _mulai_pada = time.time()
    _sumber = "sapu-belum"
    PEKERJA.clear()
    p = paralel.Pekerja(1, [], total=jumlah,
                        kerjakan=lambda h, k, c: sapu_pelayanan.kerjakan(h, k, c, jumlah, jawaban))
    p.kerja.lapor(f"Alur: Sapu Belum Pemeriksaan (tanpa berkas) · dewasa/lansia · maksimal {jumlah} peserta")
    PEKERJA.append(p)
    _kabar.update(keadaan="berjalan", pesan="")
    p.mulai()
    return {"ok": True}


def _mulai_tatalaksana(data: dict) -> dict:
    """
    Alur Tatalaksana Gigi: tanpa berkas, pesertanya diambil dari portal.

    Selalu SATU pekerja - semua browser melihat daftar yang sama, dan dua
    browser yang mengambil baris teratas bersamaan akan bertabrakan.
    """
    global PEKERJA, _mulai_pada, _sumber
    jawaban, galat = tatalaksana.urai_jawaban(data.get("tatalaksana"))
    if galat:
        return {"ok": False, "pesan": galat}
    # Kosong/0 DITOLAK, bukan diganti bawaan: angka yang tidak sengaja
    # tidak boleh diam-diam menjadi 500 peserta.
    try:
        jumlah = int(data.get("jumlah"))
    except (TypeError, ValueError):
        return {"ok": False, "pesan": "Isi jumlah peserta (angka)."}
    if not 1 <= jumlah <= config.TATALAKSANA_MAKS:
        return {"ok": False,
                "pesan": f"Jumlah peserta 1 sampai {config.TATALAKSANA_MAKS} per jalan."}
    config.TAHAP = "tatalaksana_gigi"
    _mulai_pada = time.time()
    _sumber = "tatalaksana-gigi"
    PEKERJA.clear()
    p = paralel.Pekerja(1, [], total=jumlah,
                        kerjakan=lambda h, k, c: tatalaksana.kerjakan(h, k, c, jumlah, jawaban))
    f = config.TATALAKSANA_FILTER
    p.kerja.lapor(f"Alur: Tatalaksana Gigi · {f['tahun']} · {f['kelompok']} · {f['hasil']}"
                  f" · maksimal {jumlah} peserta")
    PEKERJA.append(p)
    _kabar.update(keadaan="berjalan", pesan="")
    p.mulai()
    return {"ok": True}


def _sedang_jalan() -> bool:
    return bool(_penyiap and _penyiap.is_alive()) or any(p.hidup() for p in PEKERJA)


def _keadaan() -> dict:
    """Keadaan gabungan untuk halaman aplikasi."""
    if not PEKERJA:
        return {"keadaan": _kabar["keadaan"], "pesan": _kabar["pesan"], "total": 0,
                "selesai": 0, "hasil": [], "catatan": [], "berkas_hasil": "",
                "pekerja": [], "mulai_pada": _mulai_pada}
    r = paralel.ringkas(PEKERJA)
    r["mulai_pada"] = _mulai_pada
    # Catatan digabung dan diberi tanda pekerjanya, supaya satu layar tetap
    # bisa dibaca walau tiga browser bekerja bersamaan.
    catatan = []
    for w in r["pekerja"]:
        catatan += [f"[P{w['nomor']}] {b}" for b in w["catatan"]]
    r["catatan"] = sorted(catatan)[-160:]
    r["pesan"] = next((w["pesan"] for w in r["pekerja"] if w["pesan"]), "")
    r["berkas_hasil"] = ", ".join(w["berkas"] for w in r["pekerja"] if w["berkas"])
    return r


# ------------------------------------------------------------------- server
class Penangan(BaseHTTPRequestHandler):
    def log_message(self, *a):            # jangan penuhi terminal
        pass

    def _kirim(self, isi: bytes, jenis: str = "application/json") -> None:
        self.send_response(200)
        self.send_header("Content-Type", jenis + "; charset=utf-8")
        self.send_header("Content-Length", str(len(isi)))
        self.end_headers()
        self.wfile.write(isi)

    def _json(self, data: dict, kode: int = 200) -> None:
        isi = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(kode)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(isi)))
        self.end_headers()
        self.wfile.write(isi)

    # ------------------------------------------------------------- sesi
    def _sesi(self) -> dict | None:
        """Sesi yang sah dari cookie, atau None."""
        mentah = self.headers.get("Cookie")
        if not mentah:
            return None
        try:
            token = SimpleCookie(mentah).get("ckg_sesi")
        except Exception:
            return None
        d = SESI.get(token.value) if token else None
        if not d:
            return None
        if d["sampai"] < datetime.now():
            SESI.pop(token.value, None)
            return None
        return d

    def _wajib_masuk(self) -> bool:
        """False + jawaban 401 bila belum masuk. Dipakai semua endpoint."""
        if self._sesi():
            return True
        self._json({"ok": False, "pesan": "Sesi berakhir. Silakan masuk lagi.",
                    "perlu_masuk": True}, kode=401)
        return False

    def do_GET(self):
        jalan = urlparse(self.path).path
        if jalan in ("/", "/index.html"):
            # Belum masuk -> halaman masuk. Halaman aplikasinya sendiri tidak
            # pernah dikirim, jadi isinya tidak bisa diintip tanpa akun.
            halaman = "antarmuka.html" if self._sesi() else "masuk.html"
            self._kirim((KODE / halaman).read_bytes(), "text/html")
        elif jalan == "/panduan":
            # Panduan PDF ikut di folder aplikasi. Boleh dibuka tanpa masuk:
            # isinya cara memasang dan memakai, tidak memuat data apa pun -
            # dan justru dibutuhkan oleh yang belum bisa masuk.
            berkas = config.AKAR / "Panduan CKG Robot.pdf"
            if not berkas.exists():
                self._kirim("Panduan CKG Robot.pdf tidak ada di folder aplikasi.".encode("utf-8"),
                            "text/plain")
                return
            isi = berkas.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "application/pdf")
            self.send_header("Content-Disposition", 'inline; filename="Panduan CKG Robot.pdf"')
            self.send_header("Content-Length", str(len(isi)))
            self.end_headers()
            self.wfile.write(isi)
        elif jalan == "/admin":
            # Halaman kelola petugas - hanya untuk admin login online.
            sesi = self._sesi() or {}
            if ((sesi.get("online") or {}).get("peran")) == "admin":
                self._kirim((KODE / "admin.html").read_bytes(), "text/html")
            else:
                self.send_response(302)
                self.send_header("Location", "/")
                self.end_headers()
        elif jalan == "/api/perlu-akun":
            # Dipakai halaman masuk untuk memberi tahu bila belum ada akun,
            # dan untuk memilih tampilan masuk lokal atau online.
            if akun_online.aktif():
                self._json({"ada_akun": True, "mode": "online", "versi": config.VERSI})
            else:
                self._json({"ada_akun": akun.ada_akun(), "mode": "lokal", "versi": config.VERSI})
        elif not jalan.startswith("/api/"):
            self.send_error(404)
        elif not self._wajib_masuk():
            return
        elif jalan == "/api/awal":
            self._json({"berkas": daftar_csv(),
                        "tahap": config.TAHAP,
                        "kategori": config.KATEGORI,
                        "mode_uji": config.MODE_UJI,
                        "maks_pekerja": config.MAKS_PEKERJA,
                        "tatalaksana": tatalaksana.pilihan_untuk_aplikasi(),
                        "sapu": sapu_pelayanan.pilihan_untuk_aplikasi(),
                        "pengguna": (self._sesi() or {}).get("nama", ""),
                        "peran": ((self._sesi() or {}).get("online") or {}).get("peran", ""),
                        "pengumuman": ((self._sesi() or {}).get("online") or {}).get("pengumuman", ""),
                        "pembaruan": _pembaruan_tersedia(self._sesi()),
                        "online": bool((self._sesi() or {}).get("online")),
                        "versi": config.VERSI})
        elif jalan == "/api/perbarui-status":
            self._json(dict(_pembaruan_keadaan))
        elif jalan == "/api/admin/paket-lokal":
            if ((self._sesi() or {}).get("online") or {}).get("peran") != "admin":
                self._json({"ok": False, "pesan": "Hanya admin."}, kode=403)
            else:
                self._json({"ok": True, "paket": _paket_lokal(), "folder": str(config.PROYEK)})
        elif jalan == "/api/status":
            self._json(_keadaan())
        elif jalan == "/api/export":
            pilih = parse_qs(urlparse(self.path).query).get("status", [""])[0]
            self._export(pilih.strip())
        elif jalan == "/api/laporan":
            self._laporan()
        else:
            self.send_error(404)

    def _masuk(self, data: dict) -> None:
        """
        Periksa nama + kata sandi, lalu pasang cookie sesi.

        Ada penggemboksan sederhana: sesudah 5 kali gagal, alamat yang sama
        harus menunggu 5 menit. Tanpa itu, kata sandi pendek bisa ditebak
        ribuan kali per menit oleh siapa pun yang bisa menjangkau port ini.
        """
        alamat = self.client_address[0]
        catat = _gagal.get(alamat)
        if catat and catat[0] >= BATAS_GAGAL:
            sisa = catat[1] + JEDA_GEMBOK - datetime.now()
            if sisa.total_seconds() > 0:
                self._json({"ok": False,
                            "pesan": f"Terlalu banyak percobaan. Tunggu "
                                     f"{int(sisa.total_seconds() // 60) + 1} menit lagi."},
                           kode=429)
                return
            _gagal.pop(alamat, None)

        nama = str(data.get("nama", "")).strip()
        sandi = str(data.get("sandi", ""))
        if akun_online.aktif():
            self._masuk_online(alamat, nama, sandi)
            return
        if not akun.ada_akun():
            self._json({"ok": False,
                        "pesan": "Belum ada akun. Buat lebih dulu lewat "
                                 "buat_akun.bat di folder aplikasi."}, kode=403)
            return

        time.sleep(0.3)          # perlambat penebakan, tak terasa bagi manusia
        if not akun.periksa(nama, sandi):
            c = _gagal.setdefault(alamat, [0, datetime.now()])
            c[0] += 1
            c[1] = datetime.now()
            sisa = BATAS_GAGAL - c[0]
            self._json({"ok": False,
                        "pesan": "Nama pengguna atau kata sandi salah."
                                 + (f" Sisa {sisa} percobaan." if 0 < sisa <= 2 else "")},
                       kode=401)
            return

        _gagal.pop(alamat, None)
        self._pasang_sesi(_buat_sesi(nama.strip().lower()))

    def _pasang_sesi(self, token: str) -> None:
        isi = json.dumps({"ok": True}).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        # HttpOnly: tidak terbaca JavaScript. SameSite=Strict: tidak ikut
        # terkirim dari halaman lain.
        self.send_header("Set-Cookie",
                         f"ckg_sesi={token}; Path=/; HttpOnly; SameSite=Strict")
        self.send_header("Content-Length", str(len(isi)))
        self.end_headers()
        self.wfile.write(isi)

    def _masuk_online(self, alamat: str, email: str, sandi: str) -> None:
        """
        Masuk dengan akun Supabase, lalu tanya izinnya (disetujui admin?
        komputer ini masih dalam batas? versinya masih boleh?).
        """
        def gagal(pesan: str, kode: int = 401, hitung: bool = True) -> None:
            if hitung:
                c = _gagal.setdefault(alamat, [0, datetime.now()])
                c[0] += 1
                c[1] = datetime.now()
            self._json({"ok": False, "pesan": pesan}, kode=kode)

        if not email or not sandi:
            gagal("Isi email dan kata sandi.", 400, hitung=False)
            return
        try:
            token = akun_online.masuk(email, sandi)
        except akun_online.GagalOnline as e:
            # Galat jaringan bukan tebakan sandi - jangan ikut menggembok.
            gagal(e.pesan, 401, hitung=e.alasan not in ("jaringan",))
            return
        on = {"token": token, "email": email.strip().lower(), "peran": ""}
        izin = _izin_online({"online": on})
        if not izin.get("ok") and izin.get("alasan") == "versi_lama":
            # Versi terlalu lama: boleh MASUK supaya bisa menekan "Perbarui
            # sekarang", tetapi /api/mulai tetap menolak sampai diperbarui.
            on["peran"] = ""
            _gagal.pop(alamat, None)
            self._pasang_sesi(_buat_sesi(izin.get("nama") or email.strip().lower(), on))
            return
        if not izin.get("ok"):
            akun_online.keluar(token)
            gagal(izin.get("pesan") or "Akun belum boleh dipakai.", 403, hitung=False)
            return
        _gagal.pop(alamat, None)
        self._pasang_sesi(_buat_sesi(izin.get("nama") or email.strip().lower(), on))

    def _daftar_online(self, data: dict) -> None:
        """Pendaftaran petugas baru - statusnya 'menunggu' sampai disetujui admin."""
        if not akun_online.aktif():
            self._json({"ok": False, "pesan": "Pendaftaran hanya ada di mode login online."}, kode=400)
            return
        isian = {k: str(data.get(k, "")).strip() for k in ("email", "nama", "puskesmas", "no_hp")}
        sandi = str(data.get("sandi", ""))
        if not all(isian.values()):
            self._json({"ok": False, "pesan": "Semua isian wajib diisi."}, kode=400)
            return
        if "@" not in isian["email"]:
            self._json({"ok": False, "pesan": "Alamat email tidak sah."}, kode=400)
            return
        if len(sandi) < 8:
            self._json({"ok": False, "pesan": "Kata sandi minimal 8 karakter."}, kode=400)
            return
        try:
            hasil = akun_online.daftar(isian["email"], sandi, isian["nama"], isian["puskesmas"],
                                       isian["no_hp"])
        except akun_online.GagalOnline as e:
            self._json({"ok": False, "pesan": e.pesan}, kode=400)
            return
        if hasil["perlu_konfirmasi"]:
            self._json({"ok": True, "pesan": "Pendaftaran terkirim, tetapi server meminta konfirmasi "
                        "email. Beri tahu admin untuk mematikan 'Confirm email' di Supabase."})
            return
        self._json({"ok": True, "pesan": "Pendaftaran terkirim. Akunmu bisa dipakai setelah "
                    "disetujui admin."})

    def _admin(self, aksi: str, data: dict) -> None:
        """Kelola petugas - diteruskan ke fungsi admin_* di Supabase."""
        sesi = self._sesi()
        if aksi == "daftar":
            self._json(_rpc_admin(sesi, "admin_daftar", {}))
        elif aksi == "ubah":
            arg = {"p_id": str(data.get("id", ""))}
            for k in ("status", "peran", "nama", "puskesmas", "no_hp"):
                if data.get(k) is not None:
                    arg[f"p_{k}"] = str(data[k])
            if data.get("maks_perangkat") is not None:
                arg["p_maks_perangkat"] = int(data["maks_perangkat"])
            self._json(_rpc_admin(sesi, "admin_ubah", arg))
        elif aksi == "lepas-perangkat":
            self._json(_rpc_admin(sesi, "admin_lepas_perangkat", {"p_id": int(data.get("id") or 0)}))
        elif aksi == "ganti-sandi":
            self._json(_rpc_admin(sesi, "admin_ganti_sandi",
                                  {"p_id": str(data.get("id", "")), "p_sandi": str(data.get("sandi", ""))}))
        elif aksi == "ganti-email":
            # supabase/tambahan-ubah-email.sql
            self._json(_rpc_admin(sesi, "admin_ganti_email",
                                  {"p_id": str(data.get("id", "")), "p_email": str(data.get("email", "")).strip()}))
        elif aksi == "pemakaian":
            try:
                hari = int(data.get("hari") or 14)
            except (TypeError, ValueError):
                hari = 14
            self._json(_rpc_admin(sesi, "admin_pemakaian", {"p_hari": hari}))
        elif aksi == "pengaturan":
            arg = {}
            if data.get("versi_minimal") is not None:
                arg["p_versi_minimal"] = str(data["versi_minimal"]).strip()
            if data.get("maks_perangkat") is not None:
                arg["p_maks_perangkat"] = int(data["maks_perangkat"])
            if data.get("pengumuman") is not None:
                arg["p_pengumuman"] = str(data["pengumuman"])
            self._json(_rpc_admin(sesi, "admin_pengaturan", arg))
        elif aksi == "terbitkan":
            self._json(self._terbitkan(sesi, data))
        elif aksi == "buat-akun":
            # Admin membuatkan akun: daftar biasa, lalu langsung disetujui.
            if ((sesi or {}).get("online") or {}).get("peran") != "admin":
                self._json({"ok": False, "pesan": "Hanya admin."}, kode=403)
                return
            sandi = str(data.get("sandi", ""))
            if len(sandi) < 8:
                self._json({"ok": False, "pesan": "Kata sandi minimal 8 karakter."})
                return
            try:
                baru = akun_online.daftar(str(data.get("email", "")), sandi, str(data.get("nama", "")),
                                          str(data.get("puskesmas", "")), str(data.get("no_hp", "")))
            except akun_online.GagalOnline as e:
                self._json({"ok": False, "pesan": e.pesan})
                return
            if not baru["id"]:
                self._json({"ok": False, "pesan": "Akun dibuat, tetapi server tidak memberi ID-nya. "
                            "Setujui manual dari daftar."})
                return
            self._json(_rpc_admin(sesi, "admin_ubah", {"p_id": baru["id"], "p_status": "aktif"}))
        else:
            self._json({"ok": False, "pesan": "Aksi tidak dikenal."}, kode=404)

    def _terbitkan(self, sesi: dict | None, data: dict) -> dict:
        """
        Unggah satu ZIP buatan buat_paket.bat lalu catat versinya + sidik
        SHA-256 di Supabase. Sesudah ini aplikasi petugas menawarkan pembaruan.
        """
        on = (sesi or {}).get("online") or {}
        if on.get("peran") != "admin":
            return {"ok": False, "pesan": "Hanya admin."}
        nama = str(data.get("berkas", ""))
        pilihan = {p["nama"]: p for p in _paket_lokal()}
        if nama not in pilihan:                      # hanya nama dari daftar - bukan jalur bebas
            return {"ok": False, "pesan": "Paket tidak ditemukan di folder induk aplikasi."}
        jalur = config.PROYEK / nama
        try:
            with zipfile.ZipFile(jalur) as zf:
                versi = pembaruan.versi_dalam_zip(zf)
                lengkap = {"inti/server.py", "inti/config.py"} <= set(zf.namelist())
        except zipfile.BadZipFile:
            return {"ok": False, "pesan": "Berkas itu bukan ZIP yang sah."}
        if not versi or not lengkap:
            return {"ok": False, "pesan": "Paket tidak lengkap atau tanpa nomor versi. Buat ulang dengan "
                                          "buat_paket.bat."}
        isi = jalur.read_bytes()
        sidik = hashlib.sha256(isi).hexdigest()
        try:
            with _kunci_online:
                akun_online.unggah(on["token"], pembaruan.NAMA_PAKET, isi, "application/zip")
                hasil = akun_online.rpc("admin_terbitkan", {"p_versi": versi, "p_sha256": sidik,
                                                            "p_ukuran": len(isi),
                                                            "p_catatan": str(data.get("catatan", ""))[:1000]},
                                        on["token"])
        except akun_online.GagalOnline as e:
            return {"ok": False, "pesan": e.pesan}
        if hasil.get("ok"):
            hasil["pesan"] = (f"Versi {versi} diterbitkan ({len(isi) // 1024} KB). Petugas akan melihat "
                              "tombol Perbarui saat masuk berikutnya.")
        return hasil

    def _export(self, status: str = "") -> None:
        """
        Kirim hasil pekerjaan yang sedang tampil sebagai berkas Excel.

        `status` mempersempitnya ke satu status saja - dipakai keping rekap yang
        bisa diklik, supaya petugas bisa mengambil daftar "gagal" atau
        "tidak_ditemukan" saja tanpa menyaring sendiri di Excel.
        """
        baris = list(_keadaan()["hasil"])
        if status:
            baris = [b for b in baris if (b.get("status") or "").strip() == status]
        if not baris:
            self._json({"ok": False,
                        "pesan": (f"Tidak ada peserta berstatus '{status}'."
                                  if status else "Belum ada hasil untuk diekspor.")})
            return
        try:
            isi = buat_xlsx(baris)
        except ImportError:
            self._json({"ok": False,
                        "pesan": "Pustaka openpyxl belum terpasang. "
                                 "Jalankan pasang.bat sekali lagi."})
            return
        except Exception as e:
            self._json({"ok": False, "pesan": f"Gagal membuat berkas: {type(e).__name__}"})
            return

        # Nama berkas memuat statusnya, supaya beberapa unduhan tidak
        # saling tertukar di folder Unduhan.
        tanda = f"-{re.sub(r'[^a-z0-9]+', '-', status.lower())}" if status else ""
        nama = f"hasil{tanda}-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-"
                                         "officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{nama}"')
        self.send_header("Content-Length", str(len(isi)))
        self.end_headers()
        self.wfile.write(isi)

    def _laporan(self) -> None:
        """
        Kirim LAPORAN: berkas yang dijalankan + kolom Hasil.

        Berbeda dari `_export` yang hanya berisi hasil. Lihat
        `buat_laporan_xlsx`.
        """
        baris = list(_keadaan()["hasil"])
        if _sumber in ("tatalaksana-gigi", "sapu-belum"):
            self._json({"ok": False,
                        "pesan": "Alur ini tidak memakai berkas data - pakai "
                                 "Unduh Excel untuk hasilnya."})
            return
        if not _sumber:
            self._json({"ok": False,
                        "pesan": "Belum ada berkas yang dijalankan di sesi ini."})
            return
        if not (config.DATA / _sumber).exists():
            self._json({"ok": False,
                        "pesan": f"Berkas sumber '{_sumber}' tidak ada lagi di folder data."})
            return
        try:
            isi, tanda = buat_laporan_xlsx(_sumber, baris)
        except ImportError:
            self._json({"ok": False,
                        "pesan": "Pustaka openpyxl belum terpasang. "
                                 "Jalankan pasang.bat sekali lagi."})
            return
        except Exception as e:
            self._json({"ok": False,
                        "pesan": f"Gagal membuat laporan: {type(e).__name__}: {e}"})
            return

        nama = f"laporan-{re.sub(r'[^A-Za-z0-9]+', '-', tanda).strip('-')}-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-"
                                         "officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{nama}"')
        self.send_header("Content-Length", str(len(isi)))
        self.end_headers()
        self.wfile.write(isi)

    def do_POST(self):
        global _penyiap
        jalan = urlparse(self.path).path
        panjang = int(self.headers.get("Content-Length") or 0)
        # Berkas mentah ikut lewat sini sebagai base64. Register sekolah yang
        # paling besar pun cuma ~225 KB, jadi 40 MB sudah sangat longgar - dan
        # tanpa batas apa pun, satu permintaan salah bisa menghabiskan memori.
        if panjang > 40 * 1024 * 1024:
            self._json({"ok": False, "pesan": "Berkas terlalu besar (maksimal 40 MB)."}, kode=413)
            return
        data = json.loads(self.rfile.read(panjang) or b"{}")

        if jalan == "/api/masuk":
            self._masuk(data)
            return
        if jalan == "/api/daftar":
            self._daftar_online(data)
            return
        if jalan == "/api/keluar":
            mentah = self.headers.get("Cookie")
            if mentah:
                t = SimpleCookie(mentah).get("ckg_sesi")
                if t:
                    lama = SESI.pop(t.value, None) or {}
                    if lama.get("online"):
                        threading.Thread(target=akun_online.keluar, args=(lama["online"]["token"],),
                                         daemon=True).start()
            self._json({"ok": True})
            return
        if not self._wajib_masuk():
            return

        if jalan.startswith("/api/admin/"):
            self._admin(jalan.rsplit("/", 1)[-1], data)
            return

        if jalan == "/api/periksa-pembaruan":
            # Tanya ulang server tanpa keluar-masuk (tombol "Periksa pembaruan"
            # dan pemeriksaan otomatis tiap jam dari halaman).
            sesi = self._sesi()
            if not (sesi or {}).get("online"):
                self._json({"ok": False, "pesan": "Pemeriksaan pembaruan hanya ada di mode login online."})
                return
            izin = _izin_online(sesi)
            if not izin.get("ok") and izin.get("alasan") != "versi_lama":
                self._json({"ok": False, "pesan": izin.get("pesan") or "Tidak bisa memeriksa pembaruan."})
                return
            self._json({"ok": True, "pembaruan": _pembaruan_tersedia(sesi), "versi": config.VERSI})
            return

        if jalan == "/api/kabar":
            # Halaman bertanya tiap 2 menit - JUGA selama robot bekerja - supaya
            # pengumuman admin sampai tanpa keluar-masuk. Hanya membaca izin;
            # pekerjaan yang sedang berjalan tidak pernah dihentikan dari sini.
            sesi = self._sesi()
            on = (sesi or {}).get("online")
            if not on:
                self._json({"ok": True, "online": False})
                return
            sekarang = time.time()
            if sekarang - on.get("kabar_waktu", 0) >= KABAR_JEDA:      # beberapa tab -> satu panggilan
                on["kabar_waktu"] = sekarang
                izin = _izin_online(sesi)
                alasan = "" if izin.get("ok") else izin.get("alasan", "")
                # Gangguan jaringan/server bukan alasan memberi tahu petugas apa pun.
                if alasan in ("nonaktif", "akun_tidak_ada", "menunggu"):
                    on["kabar_akun"] = izin.get("pesan") or "Akun ini tidak lagi aktif."
                elif izin.get("ok"):
                    on["kabar_akun"] = ""
            self._json({"ok": True, "online": True, "pengumuman": on.get("pengumuman", ""),
                        "akun": on.get("kabar_akun", "")})
            return

        if jalan == "/api/perbarui":
            sesi = self._sesi()
            if not _pembaruan_tersedia(sesi):
                self._json({"ok": False, "pesan": "Tidak ada pembaruan untuk versi ini."})
                return
            if _sedang_jalan():
                self._json({"ok": False, "pesan": "Robot sedang bekerja. Perbarui setelah pekerjaan selesai "
                                                  "atau dihentikan."})
                return
            if not _kunci_pembaruan.acquire(blocking=False):
                self._json({"ok": False, "pesan": "Pembaruan sedang berjalan."})
                return
            _pembaruan_keadaan.clear()
            _pembaruan_keadaan.update(tahap="mulai", pesan="Menyiapkan…", persen=1, galat="")

            def kerja():
                try:
                    _kerjakan_pembaruan(sesi)
                finally:
                    if _pembaruan_keadaan.get("tahap") != "mulai_ulang":
                        _kunci_pembaruan.release()
            threading.Thread(target=kerja, daemon=True).start()
            self._json({"ok": True})
            return

        if jalan == "/api/konversi":
            self._json(_konversi(data))
            return

        if jalan == "/api/mulai":
            # Login online: akun ditanya ulang setiap kali pekerjaan dimulai.
            izin = _izin_online(self._sesi())
            if not izin.get("ok"):
                self._json({"ok": False, "pesan": izin.get("pesan") or "Akun tidak boleh dipakai.",
                            "perlu_perbarui": izin.get("alasan") == "versi_lama"}, kode=403)
                return
            # Hitungan pekerjaan sebelumnya dikirim dulu, sebelum daftar
            # pekerjanya dikosongkan (lihat _kirim_pemakaian).
            try:
                _kirim_pemakaian()
            except Exception:
                pass
            global _sesi_kerja
            _sesi_kerja = self._sesi()
            if _sedang_jalan() and not data.get("paksa"):
                hidup = [p.nomor for p in PEKERJA if p.hidup()]
                self._json({"ok": False, "bisa_paksa": True,
                            "pesan": "Masih ada pekerjaan sebelumnya yang berjalan"
                                     + (f" (pekerja {hidup})" if hidup else "")
                                     + ". Tutup jendela browsernya lalu tekan "
                                       "“Mulai paksa” bila memang mau memulai baru."})
                return
            if data.get("paksa"):
                # Pekerja lama DITINGGALKAN begitu saja - utasnya daemon, jadi
                # ikut mati saat aplikasi ditutup.
                #
                # Kuncinya JANGAN dilepas. Percobaan sebelumnya melepas
                # `tanda_login` sekaligus menyalakan `minta_berhenti`: pekerja
                # lama jadi "bangun", membuat berkas hasil kosong, lalu berhenti
                # di pemeriksaan pertama - meninggalkan dua berkas berisi baris
                # judul saja yang membingungkan saat dibaca nanti.
                for p in PEKERJA:
                    p.kerja.minta_berhenti = True
                    p.kerja.lapor("Ditinggalkan karena pekerjaan baru dimulai paksa.")
                PEKERJA.clear()

            # Pilihan di halaman MENGGANTIKAN pengaturan.txt selama sesi ini.
            # Nilai yang tidak dikenal diabaikan, bukan dipaksakan: salah ketik
            # tidak boleh diam-diam menjalankan alur yang lain.
            jenis = str(data.get("jenis", "")).strip().lower()
            alur = str(data.get("alur", "")).strip().lower()
            if jenis in ("auto", "dewasa", "sekolah", "balita"):
                config.KATEGORI = jenis
            if alur == "sapu_belum":
                self._json(_mulai_sapu(data))
                return
            if alur == "tatalaksana_gigi":
                self._json(_mulai_tatalaksana(data))
                return
            if alur in ("penuh", "pelayanan", "daftar_pelayanan",
                        "pelayanan_sapu", "daftar_konfirmasi"):
                config.TAHAP = alur
            try:
                jumlah = max(1, min(int(data.get("pekerja", 1)),
                                    config.MAKS_PEKERJA))
            except (TypeError, ValueError):
                jumlah = 1

            PEKERJA.clear()
            _kabar.update(keadaan="berjalan", pesan="")
            _penyiap = threading.Thread(
                target=_siapkan_dan_jalankan, daemon=True,
                args=(data.get("berkas", ""), data.get("rentang", ""), jumlah))
            _penyiap.start()
            self._json({"ok": True})

        elif jalan == "/api/sudah-login":
            # Tiap pekerja login sendiri-sendiri; nomor 0 berarti semuanya.
            nomor = int(data.get("pekerja", 0) or 0)
            for p in PEKERJA:
                if nomor in (0, p.nomor):
                    p.kerja.tanda_login.set()
            self._json({"ok": True})

        elif jalan == "/api/fokus":
            # Memunculkan jendela browser pekerja. TIDAK lewat Playwright:
            # objeknya terikat pada utas pekerja, dan menyentuhnya dari utas
            # server ini bisa merusak sesi yang sedang bekerja.
            nomor = int(data.get("pekerja", 0) or 0)
            sasaran = next((p for p in PEKERJA if p.nomor == nomor), None)
            if sasaran is None:
                self._json({"ok": False, "pesan": f"Pekerja {nomor} tidak ada."})
                return
            galat = jendela.munculkan(sasaran.profil)
            self._json({"ok": not galat, "pesan": galat})

        elif jalan == "/api/berhenti":
            for p in PEKERJA:
                p.kerja.minta_berhenti = True
                p.kerja.lapor("Permintaan berhenti diterima — berhenti setelah peserta ini.")
            self._json({"ok": True})

        else:
            self.send_error(404)


def _catat_galat(judul: str, galat: BaseException) -> None:
    """
    Tulis galat ke berkas, bukan hanya ke jendela hitam.

    Jendela hitam bisa tertutup, tergulir, atau tidak diperhatikan - dan saat
    aplikasi berhenti menjawab, justru itulah satu-satunya keterangan yang ada.
    """
    import traceback
    try:
        config.siapkan_folder()
        with (config.HASIL / "galat-aplikasi.txt").open("a", encoding="utf-8") as f:
            f.write(f"\n===== {datetime.now():%Y-%m-%d %H:%M:%S} · {judul} =====\n")
            f.write("".join(traceback.format_exception(type(galat), galat,
                                                       galat.__traceback__)))
    except Exception:
        pass


def _galat_utas(info) -> None:
    _catat_galat(f"utas {info.thread.name if info.thread else '?'}", info.exc_value)


threading.excepthook = _galat_utas


class _ServerIPv6(ThreadingHTTPServer):
    # Lihat `_sudah_dipakai`: tanpa ini Windows mengizinkan server kedua
    # mengikat alamat yang sama diam-diam.
    allow_reuse_address = False

    address_family = socket.AF_INET6


def _sudah_dipakai() -> bool:
    """
    Benar bila sudah ada aplikasi yang mendengarkan di port ini.

    KENAPA HARUS DIPERIKSA SENDIRI. `ThreadingHTTPServer` mewarisi
    `allow_reuse_address = 1`, dan di Windows itu berarti proses KEDUA boleh
    mengikat alamat yang SAMA tanpa galat sedikit pun. Yang terjadi kemudian:
    dua server hidup berdampingan, permintaan dilayani entah yang mana - dan
    biasanya yang LAMA, karena ia lebih dulu memegang 127.0.0.1.

    Akibatnya membingungkan dan sulit ditebak: kode yang baru saja diperbaiki
    seolah tidak berpengaruh apa-apa, sebab yang menjawab masih proses lama
    yang memuat modul versi sebelumnya. Persis itu yang terjadi 2 Sep 2026 -
    proses pukul 13:06 masih memegang 127.0.0.1:8765 saat proses pukul 14:13
    dijalankan, dan pengujian pukul 14:17 memakai kode lama.
    """
    for keluarga, alamat in ((socket.AF_INET, ("127.0.0.1", PORTA)),
                             (socket.AF_INET6, ("::1", PORTA))):
        try:
            with socket.socket(keluarga, socket.SOCK_STREAM) as s:
                s.settimeout(0.6)
                if s.connect_ex(alamat) == 0:
                    return True
        except OSError:
            continue
    return False


def _layani() -> None:
    """
    Dengarkan di KEDUA alamat loopback: 127.0.0.1 dan ::1.

    Di banyak komputer Windows, nama `localhost` menunjuk ::1 lebih dulu.
    Kalau server hanya mendengarkan IPv4, tiap permintaan mencoba IPv6,
    MENGGANTUNG sekitar 2 detik, baru mundur ke IPv4 - dan halaman yang
    menanya tiap 1,5 detik jadi menumpuk sampai tampak membeku.

    Keduanya tetap alamat loopback, jadi aplikasi ini tetap hanya bisa
    dibuka dari komputer ini sendiri.
    """
    class _ServerIPv4(ThreadingHTTPServer):
        allow_reuse_address = False       # lihat `_sudah_dipakai`

    utama = _ServerIPv4(("127.0.0.1", PORTA), Penangan)
    _SERVER.append(utama)
    try:
        enam = _ServerIPv6(("::1", PORTA), Penangan)
        _SERVER.append(enam)
    except OSError as e:          # komputer tanpa IPv6 - tidak apa-apa
        enam = None
        print(f"  (IPv6 dilewati: {e.strerror or e})")
    if enam is not None:
        threading.Thread(target=enam.serve_forever, daemon=True).start()
    utama.serve_forever()
    for s in _SERVER:             # dilepas supaya server versi baru bisa memakai port ini
        s.server_close()


def _atur_jendela(kecilkan: bool) -> None:
    """
    Kecilkan jendela hitam ke taskbar setelah aplikasi siap (atau munculkan
    lagi bila ada galat). Hanya bila dibuka lewat aplikasi.bat (CKG_KECILKAN):
    yang menjalankan server.py langsung dari terminal tidak ikut dikecilkan.
    """
    if os.name != "nt" or not os.environ.get("CKG_KECILKAN"):
        return
    try:
        import ctypes
        h = ctypes.windll.kernel32.GetConsoleWindow()
        if h:
            ctypes.windll.user32.ShowWindow(h, 6 if kecilkan else 9)   # SW_MINIMIZE / SW_RESTORE
    except Exception:
        pass


def main() -> None:
    alamat = f"http://localhost:{PORTA}"
    print("=" * 60)
    print("  CKG Dewasa/Lansia — antarmuka aplikasi")
    print(f"  Buka di browser: {alamat}")
    print("  Tutup jendela ini untuk mematikan aplikasi.")
    print("=" * 60)
    # Diperiksa SEBELUM apa pun dibuka. Menjalankan aplikasi dua kali adalah
    # kekeliruan yang gampang terjadi (jendela lama belum ditutup), dan tanpa
    # penjaga ini akibatnya tidak terlihat sama sekali - lihat `_sudah_dipakai`.
    if _sudah_dipakai():
        print(f"\n[BERHENTI] Aplikasi sudah berjalan dan memakai port {PORTA}.")
        print("  Tutup dulu jendela hitam yang lama, baru jalankan lagi.")
        print(f"  Kalau memang ingin memakai yang sedang berjalan, buka {alamat}.")
        print("\n  PENTING: selama jendela lama masih terbuka, ia memakai kode")
        print("  versi lama - perbaikan yang baru dipasang TIDAK akan berlaku.")
        # Paling sering: ikon CKG Robot diklik lagi padahal aplikasinya masih
        # jalan (jendela hitamnya mengecil di taskbar). Halamannya dibukakan.
        try:
            webbrowser.open(alamat)
        except Exception:
            pass
        input("\nTekan Enter untuk menutup...")
        return
    # Dimulai ulang sesudah pembaruan: halaman yang sudah terbuka memuat ulang
    # sendiri, jadi tidak perlu membuka tab baru.
    if not os.environ.get("CKG_MULAI_ULANG"):
        try:
            webbrowser.open(alamat)
        except Exception:
            pass
    threading.Thread(target=_pengirim_pemakaian, daemon=True).start()
    threading.Thread(target=pintasan.pastikan_sekali, daemon=True).start()
    _atur_jendela(kecilkan=True)
    try:
        _layani()
        if _mulai_ulang:
            # MULAI ULANG SESUDAH PEMBARUAN. Server versi baru dijalankan sebagai
            # anak proses ini, di jendela yang SAMA, dan proses ini menunggunya.
            # Bukan dengan keluar lalu membuka jendela baru: aplikasi.bat masih
            # menunggu proses ini, dan begitu ia berakhir jendela lama akan
            # menampilkan "Aplikasi berhenti" lalu diam menunggu tombol - dua
            # jendela hitam yang membingungkan petugas.
            print("\n" + "=" * 60)
            print("  Pembaruan terpasang - memulai ulang aplikasi versi baru…")
            print("=" * 60 + "\n")
            subprocess.call([sys.executable, str(KODE / "server.py")], cwd=str(config.AKAR),
                            env={**os.environ, "CKG_MULAI_ULANG": "1"})
            return
    except KeyboardInterrupt:
        print("\nDihentikan.")
    except Exception as e:
        _catat_galat("server berhenti", e)
        _atur_jendela(kecilkan=False)
        print(f"\n[GAGAL] {type(e).__name__}: {e}")
        print(f"Keterangan lengkap ditulis ke hasil/galat-aplikasi.txt")
        input("Tekan Enter untuk menutup...")
        raise


if __name__ == "__main__":
    main()
